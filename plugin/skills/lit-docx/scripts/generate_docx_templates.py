#!/usr/bin/env python3
"""Generate publisher-style DOCX templates from templates/registry.yaml.

Data-driven: every publisher defined in the registry becomes
``templates/docx/<publisher>.docx``. Safe to re-run; output is idempotent in
content (docx zip metadata may differ across runs, but word/document.xml and
word/styles.xml content are stable).

Usage:
    python scripts/generate_docx_templates.py
    python scripts/generate_docx_templates.py --publisher elsevier
    python scripts/generate_docx_templates.py --registry path/to/registry.yaml --out-dir path/to/docx
"""
from __future__ import annotations

import os as _os
import sys as _sys

_sys.dont_write_bytecode = True  # the installed skill directory stays read-only
_sys.path.insert(0, _os.path.join(_os.path.dirname(_os.path.abspath(__file__)), "..", "..", "lit-pptx", "scripts"))
from office_runtime_bootstrap import ensure_runtime  # noqa: E402

ensure_runtime(["docx", "yaml"])

import argparse
import sys
from pathlib import Path
from typing import Any

try:
    import yaml
except ImportError:
    print("Error: 'pyyaml' not found. pip install pyyaml", file=sys.stderr)
    raise

try:
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:
    print("Error: 'python-docx' not found. pip install python-docx", file=sys.stderr)
    raise


# ---------------------------------------------------------------------------
# OOXML helpers (python-docx high-level API does not cover these)
# ---------------------------------------------------------------------------


def set_fonts_on_rpr(rpr, latin: str, cjk: str | None) -> None:
    """Apply w:rFonts (ascii/hAnsi/cs + optional eastAsia) to an rPr element."""
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:cs"), latin)
    if cjk:
        rfonts.set(qn("w:eastAsia"), cjk)


def set_style_fonts(style, latin: str, cjk: str | None) -> None:
    """Set ascii+hAnsi+cs+eastAsia on a style's rPr."""
    element = style.element
    rpr = element.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        element.append(rpr)
    set_fonts_on_rpr(rpr, latin, cjk)


def set_run_fonts(run, latin: str, cjk: str | None) -> None:
    """Set ascii+hAnsi+cs+eastAsia on a run's rPr."""
    rpr = run._element.get_or_add_rPr()
    set_fonts_on_rpr(rpr, latin, cjk)


def ensure_ppr(element):
    ppr = element.find(qn("w:pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        element.insert(0, ppr)
    return ppr


def set_style_widow_control(style) -> None:
    ppr = ensure_ppr(style.element)
    wc = ppr.find(qn("w:widowControl"))
    if wc is None:
        wc = OxmlElement("w:widowControl")
        ppr.append(wc)
    wc.set(qn("w:val"), "1")


def set_style_hanging_indent(style, hanging_cm: float) -> None:
    """Set hanging indent on a paragraph style (twips)."""
    ppr = ensure_ppr(style.element)
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    # hanging requires left >= hanging for Word to render properly
    twips = int(round(hanging_cm * 567))
    ind.set(qn("w:left"), str(twips))
    ind.set(qn("w:hanging"), str(twips))


def enable_auto_hyphenation(doc) -> None:
    settings = doc.settings.element
    if settings.find(qn("w:autoHyphenation")) is None:
        settings.append(OxmlElement("w:autoHyphenation"))


def set_section_page(section, page_cfg: dict, title_pg: bool) -> None:
    """Set page size (A4) and margins from the registry.

    python-docx exposes Section.page_width/page_height/margins. We use those
    for the high-level attributes and drop in w:titlePg if first_page_different.
    """
    size = (page_cfg.get("size") or "A4").upper()
    if size == "A4":
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
    elif size == "LETTER":
        section.page_width = Cm(21.59)
        section.page_height = Cm(27.94)
    section.top_margin = Cm(page_cfg.get("margin_top_cm", 2.5))
    section.bottom_margin = Cm(page_cfg.get("margin_bottom_cm", 2.5))
    section.left_margin = Cm(page_cfg.get("margin_left_cm", 2.5))
    section.right_margin = Cm(page_cfg.get("margin_right_cm", 2.5))
    if title_pg:
        sect_pr = section._sectPr
        if sect_pr.find(qn("w:titlePg")) is None:
            sect_pr.append(OxmlElement("w:titlePg"))


# ---------------------------------------------------------------------------
# Style builders
# ---------------------------------------------------------------------------


def configure_normal_style(doc, publisher: dict) -> int:
    """Configure Normal style: body font, CJK pairing, indent/justify/line-spacing."""
    body = publisher["docx"]["font"]["body"]
    body_cjk = publisher["docx"]["font"].get("body_cjk") or {}
    design = publisher["design"]
    para_cfg = design["paragraph"]

    style = doc.styles["Normal"]
    style.font.name = body["family"]
    style.font.size = Pt(body["size_pt"])
    set_style_fonts(style, body["family"], body_cjk.get("family"))

    pf = style.paragraph_format
    indent_cm = para_cfg.get("first_line_indent_cm", 0)
    if indent_cm:
        pf.first_line_indent = Cm(indent_cm)
    line_spacing = para_cfg.get("line_spacing") or body.get("line_spacing") or 1.15
    pf.line_spacing = float(line_spacing)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.space_before = Pt(para_cfg.get("space_before_pt", 0))
    pf.space_after = Pt(para_cfg.get("space_after_pt", 0))
    if para_cfg.get("alignment") == "justify":
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif para_cfg.get("alignment") == "left":
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if para_cfg.get("widow_control"):
        set_style_widow_control(style)
    return 1


def suppress_style_numbering(style) -> None:
    """Force-disable any list or outline numbering on a paragraph style.

    Some readers (LibreOffice in particular) auto-apply a bullet / chapter
    marker to paragraphs whose style carries `w:outlineLvl` unless an explicit
    `w:numPr` with `numId=0` is set. `numId=0` is the OOXML sentinel for
    "no numbering". Set it on both ilvl and numId to override inheritance from
    any ancestor style.
    """
    ppr = style.element.find(qn("w:pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        style.element.append(ppr)
    # Remove any pre-existing numPr first (idempotent regenerate)
    for existing in ppr.findall(qn("w:numPr")):
        ppr.remove(existing)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "0")
    num_pr.append(ilvl)
    num_pr.append(num_id)
    ppr.append(num_pr)


def configure_heading_styles(doc, publisher: dict) -> int:
    """Override Heading 1..4 with publisher hierarchy values."""
    body = publisher["docx"]["font"]["body"]
    body_cjk = publisher["docx"]["font"].get("body_cjk") or {}
    latin = body["family"]
    cjk = body_cjk.get("family")
    hh = publisher["design"]["heading_hierarchy"]
    touched = 0
    for level, key in enumerate(["h1", "h2", "h3", "h4"], start=1):
        cfg = hh.get(key)
        if cfg is None:
            continue
        style = doc.styles[f"Heading {level}"]
        style.font.name = latin
        style.font.size = Pt(cfg["size_pt"])
        weight = cfg.get("weight", "regular")
        style.font.bold = weight in ("bold", "semibold")
        style.font.italic = bool(cfg.get("italic", False))
        # All headings black (color discipline)
        style.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        set_style_fonts(style, latin, cjk)

        pf = style.paragraph_format
        pf.space_before = Pt(cfg.get("space_before_pt", 12))
        pf.space_after = Pt(cfg.get("space_after_pt", 4))
        pf.keep_with_next = bool(cfg.get("keep_with_next", True))
        pf.keep_together = True

        # Hard-disable any inherited list/outline numbering so readers never
        # render a bullet or chapter marker in front of the heading text.
        suppress_style_numbering(style)
        touched += 1
    return touched


def configure_hyperlink_style(doc, publisher: dict) -> int:
    """Override (or create) Hyperlink character style with publisher link color."""
    link_hex = publisher["design"]["color_palette"]["link"].lstrip("#")
    # On a blank Document, 'Hyperlink' char style exists after the first
    # hyperlink is inserted. We ensure it exists by creating it defensively.
    try:
        style = doc.styles["Hyperlink"]
    except KeyError:
        style = doc.styles.add_style("Hyperlink", WD_STYLE_TYPE.CHARACTER)
    style.font.color.rgb = RGBColor.from_string(link_hex)
    style.font.underline = True
    # Also set via OOXML so it persists on all readers
    rpr = style.element.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        style.element.append(rpr)
    color = rpr.find(qn("w:color"))
    if color is None:
        color = OxmlElement("w:color")
        rpr.append(color)
    color.set(qn("w:val"), link_hex)
    return 1


def configure_bibliography_style(doc, publisher: dict) -> int:
    """Create a 'Bibliography' paragraph style with hanging indent."""
    body = publisher["docx"]["font"]["body"]
    body_cjk = publisher["docx"]["font"].get("body_cjk") or {}
    bib = publisher["design"]["bibliography"]
    if "Bibliography" in [s.name for s in doc.styles]:
        style = doc.styles["Bibliography"]
    else:
        style = doc.styles.add_style("Bibliography", WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles["Normal"]
    style.font.name = body["family"]
    style.font.size = Pt(bib.get("font_size_pt", 10))
    set_style_fonts(style, body["family"], body_cjk.get("family"))
    pf = style.paragraph_format
    pf.line_spacing = float(bib.get("line_spacing", 1.15))
    pf.space_after = Pt(bib.get("space_between_entries_pt", 6))
    pf.first_line_indent = None  # hanging indent instead
    set_style_hanging_indent(style, bib.get("hanging_indent_cm", 0.6))
    return 1


def configure_caption_style(doc, publisher: dict) -> int:
    body = publisher["docx"]["font"]["body"]
    body_cjk = publisher["docx"]["font"].get("body_cjk") or {}
    # python-docx's blank document does include a 'Caption' style; if not, make it.
    try:
        style = doc.styles["Caption"]
    except KeyError:
        style = doc.styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = body["family"]
    style.font.size = Pt(publisher["design"]["figure_style"].get("caption_font_size_pt", 10))
    style.font.italic = True
    style.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    set_style_fonts(style, body["family"], body_cjk.get("family"))
    return 1


# ---------------------------------------------------------------------------
# Self-describing title block sample
# ---------------------------------------------------------------------------


def add_sample_title_block(doc, publisher: dict) -> None:
    """Add title/authors/affiliations/abstract/keywords placeholders."""
    design = publisher["design"]
    tb = design["title_block"]
    body = publisher["docx"]["font"]["body"]
    body_cjk = publisher["docx"]["font"].get("body_cjk") or {}
    latin = body["family"]
    cjk = body_cjk.get("family")

    def _p(text: str, *, size: int, bold: bool = False, italic: bool = False,
           align=WD_ALIGN_PARAGRAPH.CENTER, space_after: int = 4) -> None:
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(space_after)
        run = p.add_run(text)
        run.font.name = latin
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
        set_run_fonts(run, latin, cjk)

    _p("Manuscript Title — replace with your title",
       size=tb["title"]["size_pt"], bold=True,
       space_after=tb["title"].get("space_after_pt", 12))

    # Authors paragraph with superscript affiliation markers
    authors_p = doc.add_paragraph()
    authors_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    authors_p.paragraph_format.first_line_indent = Cm(0)
    authors_p.paragraph_format.space_after = Pt(tb["authors"].get("space_after_pt", 4))
    sep = tb.get("author_separator", ", ")
    samples = [("Author One", "1", True), ("Author Two", "2", False)]
    for idx, (name, aff, corr) in enumerate(samples):
        if idx > 0:
            sep_run = authors_p.add_run(sep)
            sep_run.font.name = latin
            sep_run.font.size = Pt(tb["authors"]["size_pt"])
            set_run_fonts(sep_run, latin, cjk)
        name_run = authors_p.add_run(name)
        name_run.font.name = latin
        name_run.font.size = Pt(tb["authors"]["size_pt"])
        set_run_fonts(name_run, latin, cjk)
        sup_run = authors_p.add_run(aff)
        sup_run.font.superscript = True
        sup_run.font.name = latin
        sup_run.font.size = Pt(tb["authors"]["size_pt"])
        set_run_fonts(sup_run, latin, cjk)
        if corr:
            star = authors_p.add_run(tb.get("corresponding_marker", "*"))
            star.font.superscript = True
            star.font.name = latin
            set_run_fonts(star, latin, cjk)

    _p("¹ Affiliation one — department, institution, city, country",
       size=tb["affiliations"]["size_pt"], italic=True, space_after=2)
    _p("² Affiliation two — department, institution, city, country",
       size=tb["affiliations"]["size_pt"], italic=True,
       space_after=tb["affiliations"].get("space_after_pt", 8))

    # Abstract label + placeholder
    ap = doc.add_paragraph()
    ap.paragraph_format.first_line_indent = Cm(0)
    ap.paragraph_format.space_after = Pt(4)
    lab = ap.add_run(tb.get("abstract_label", "Abstract"))
    lab.bold = True
    lab.font.name = latin
    lab.font.size = Pt(tb.get("abstract_font_size_pt", 10))
    set_run_fonts(lab, latin, cjk)

    abs_p = doc.add_paragraph()
    abs_p.paragraph_format.first_line_indent = Cm(0)
    abs_p.paragraph_format.space_after = Pt(6)
    abs_run = abs_p.add_run(
        "Replace this paragraph with your abstract. Keep it focused on the "
        "problem, approach, key results, and implications. Straight quotes, "
        "hyphens-as-dashes, and missing nbsp before units will be auto-fixed "
        "by convert_md_to_docx.py when --publisher is set."
    )
    abs_run.font.name = latin
    abs_run.font.size = Pt(tb.get("abstract_font_size_pt", 10))
    set_run_fonts(abs_run, latin, cjk)

    kw_p = doc.add_paragraph()
    kw_p.paragraph_format.first_line_indent = Cm(0)
    kw_p.paragraph_format.space_after = Pt(12)
    lab2 = kw_p.add_run(tb.get("keywords_label", "Keywords") + ": ")
    lab2.bold = True
    lab2.font.name = latin
    lab2.font.size = Pt(tb.get("abstract_font_size_pt", 10))
    set_run_fonts(lab2, latin, cjk)
    kw_val = kw_p.add_run("keyword1; keyword2; keyword3")
    kw_val.font.name = latin
    kw_val.font.size = Pt(tb.get("abstract_font_size_pt", 10))
    set_run_fonts(kw_val, latin, cjk)


# ---------------------------------------------------------------------------
# Main build
# ---------------------------------------------------------------------------


def build_template(publisher_key: str, publisher: dict, out_path: Path) -> int:
    doc = Document()

    # Page geometry (first section)
    section = doc.sections[0]
    page_cfg = publisher["docx"]["page"]
    hf = publisher["design"]["header_footer"]
    set_section_page(section, page_cfg, title_pg=bool(hf.get("first_page_different")))

    # Auto-hyphenation
    enable_auto_hyphenation(doc)

    # Style configuration
    style_count = 0
    style_count += configure_normal_style(doc, publisher)
    style_count += configure_heading_styles(doc, publisher)
    style_count += configure_hyperlink_style(doc, publisher)
    style_count += configure_bibliography_style(doc, publisher)
    style_count += configure_caption_style(doc, publisher)

    # Self-describing sample title block
    add_sample_title_block(doc, publisher)

    # Sample section heading so Heading 1 style is exercised
    sh = doc.add_paragraph("1 Introduction", style="Heading 1")
    sh_body = doc.add_paragraph(
        "Replace this paragraph with your manuscript body. This template was "
        "generated from templates/registry.yaml — do not hand-edit. To update "
        "style, edit the registry and rerun scripts/generate_docx_templates.py."
    )
    sh_body.paragraph_format.first_line_indent = Cm(
        publisher["design"]["paragraph"].get("first_line_indent_cm", 0.5)
    )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return style_count


def load_registry(path: Path) -> dict:
    with path.open("r", encoding="utf-8") as fh:
        data = yaml.safe_load(fh)
    if "publishers" not in data:
        raise SystemExit(f"registry {path} missing top-level 'publishers' key")
    return data


def parse_args() -> argparse.Namespace:
    script_dir = Path(__file__).resolve().parent
    default_registry = script_dir.parent / "templates" / "registry.yaml"
    default_out = script_dir.parent / "templates" / "docx"

    p = argparse.ArgumentParser(description="Generate publisher DOCX templates")
    p.add_argument("--publisher", help="Only generate this publisher (default: all)")
    p.add_argument("--registry", default=str(default_registry),
                   help=f"Path to registry.yaml (default: {default_registry})")
    p.add_argument("--out-dir", default=str(default_out),
                   help=f"Output directory for .docx files (default: {default_out})")
    return p.parse_args()


def main() -> int:
    args = parse_args()
    registry_path = Path(args.registry).expanduser().resolve()
    out_dir = Path(args.out_dir).expanduser().resolve()

    if not registry_path.exists():
        print(f"Error: registry not found: {registry_path}", file=sys.stderr)
        return 1

    registry = load_registry(registry_path)
    publishers = registry["publishers"]

    if args.publisher:
        if args.publisher not in publishers:
            print(f"Error: publisher {args.publisher!r} not in registry "
                  f"(known: {', '.join(publishers)})", file=sys.stderr)
            return 1
        targets = {args.publisher: publishers[args.publisher]}
    else:
        targets = publishers

    for key, publisher in targets.items():
        out_path = out_dir / f"{key}.docx"
        n = build_template(key, publisher, out_path)
        rel = out_path
        try:
            rel = out_path.relative_to(Path.cwd())
        except ValueError:
            pass
        print(f"generated {rel} ({n} styles applied)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
