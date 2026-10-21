#!/usr/bin/env python3
from __future__ import annotations

import os as _os
import sys as _sys

_sys.dont_write_bytecode = True  # the installed skill directory stays read-only
_sys.path.insert(0, _os.path.join(_os.path.dirname(_os.path.abspath(__file__)), "..", "..", "lit-pptx", "scripts"))
from office_runtime_bootstrap import ensure_runtime  # noqa: E402

ensure_runtime(["docx", "markdown", "bs4", "yaml"])

import argparse
import re
import sys
import zipfile
from dataclasses import dataclass
from pathlib import Path

from convert_md_to_docx import (
    detect_locale,
    fix_dashes,
    fix_ellipsis,
    fix_quotes,
    fix_unit_spacing,
    load_publisher,
    normalize_double_spaces,
    parse_frontmatter,
)


def _load_yaml():
    try:
        import yaml
        return yaml
    except ImportError:
        print("Error: pyyaml is required for slop_lint.py. pip install pyyaml", file=sys.stderr)
        raise SystemExit(2)


@dataclass
class Finding:
    rule_id: str
    line: int
    column: int
    section: str
    message: str
    excerpt: str


def newline_offsets(text: str) -> list[int]:
    offsets = [0]
    for idx, ch in enumerate(text):
        if ch == "\n":
            offsets.append(idx + 1)
    return offsets


def line_col(text: str, offsets: list[int], pos: int) -> tuple[int, int]:
    line = 1
    for idx, start in enumerate(offsets):
        if start > pos:
            break
        line = idx + 1
    col = pos - offsets[line - 1] + 1
    return line, col


def load_phrase_rules(path: Path) -> dict:
    yaml = _load_yaml()
    with path.open("r", encoding="utf-8") as fh:
        return yaml.safe_load(fh) or {}


def split_sections(body: str) -> list[tuple[str, str, int]]:
    sections: list[tuple[str, str, int]] = []
    current_heading = "Preamble"
    current_lines: list[str] = []
    current_start = 1
    line_no = 1
    for line in body.splitlines():
        if line.startswith("# "):
            if current_lines:
                sections.append((current_heading, "\n".join(current_lines).strip(), current_start))
            current_heading = line[2:].strip()
            current_lines = []
            current_start = line_no + 1
        else:
            current_lines.append(line)
        line_no += 1
    if current_lines:
        sections.append((current_heading, "\n".join(current_lines).strip(), current_start))
    return sections


def sentence_split(text: str) -> list[str]:
    items = re.split(r"(?<=[.!?])\s+", text.strip())
    return [item for item in items if item]


def word_count(text: str) -> int:
    return len(re.findall(r"\b[\w’'-]+\b", text))


TITLE_MINOR = {"a", "an", "the", "and", "but", "or", "nor", "for", "of", "in", "on", "at", "to", "by", "as", "via", "with", "from"}


def heading_case_ok(heading: str, expected: str) -> bool:
    bare = heading.strip()
    if not bare:
        return True
    if re.search(r"[가-힣]", bare):
        return True
    if expected == "sentence_case":
        return bare[0].isupper() and bare[1:] != bare[1:].upper()
    if expected == "title_case":
        # Title case capitalises every word except a short article, conjunction or preposition after the first.
        words = [w for w in re.findall(r"[A-Za-z][A-Za-z'-]*", bare)][:6]
        return all(word[0].isupper() or (i and word in TITLE_MINOR) for i, word in enumerate(words)) if words else True
    return True


def normalize_heading(value: str) -> str:
    lowered = value.strip().lower().replace("&", "and")
    lowered = re.sub(r"[^a-z0-9가-힣\s]", "", lowered)
    lowered = re.sub(r"\s+", " ", lowered)
    aliases = {
        "materials and methods": "methods",
        "methodology": "methods",
        "conclusion": "conclusions",
    }
    return aliases.get(lowered, lowered)


def passive_ratio(sentences: list[str]) -> float:
    if not sentences:
        return 0.0
    hits = 0
    for sentence in sentences:
        if re.search(r"\b(was|were|is|are|been|being)\b\s+\b\w+(ed|en)\b", sentence, re.IGNORECASE):
            hits += 1
    return hits / len(sentences)


def citation_count(text: str) -> int:
    return len(re.findall(r"\[(\d+)\]|\([A-Z][A-Za-z]+ et al\.,? \d{4}\)|\[@[^\]]+\]", text))


def type_token_ratio(text: str) -> float:
    words = [w.lower() for w in re.findall(r"\b[\w’'-]+\b", text)]
    if not words:
        return 1.0
    return len(set(words)) / len(words)


def add_finding(findings: list[Finding], text: str, offsets: list[int], pos: int, rule_id: str, section: str, message: str, excerpt: str) -> None:
    line, column = line_col(text, offsets, pos)
    findings.append(Finding(rule_id, line, column, section, message, excerpt.strip()))


def lint_text(md_text: str, publisher: dict, phrase_rules: dict, locale: str) -> tuple[list[Finding], dict, str, str]:
    frontmatter, body = parse_frontmatter(md_text)
    detected_locale = detect_locale(body) if locale == "auto" else locale
    lang_rules = phrase_rules.get("en", {})
    ko_rules = phrase_rules.get("ko", {})
    full_text = body
    offsets = newline_offsets(full_text)
    findings: list[Finding] = []
    sections = split_sections(body)

    ai_phrases = [re.compile(re.escape(item), re.IGNORECASE) for item in lang_rules.get("ai_tell_phrases", [])]
    stock_openers = [item.lower() for item in lang_rules.get("stock_openers", [])]
    stock_closers = [item.lower() for item in lang_rules.get("stock_closers", [])]
    hedging_words = set(lang_rules.get("hedging_words", []))
    redundant = set(lang_rules.get("redundant_qualifiers", []))
    adjective_stacks = [re.compile(re.escape(item), re.IGNORECASE) for item in lang_rules.get("adjective_stacks", [])]
    expected_order = publisher.get("docx", {}).get("section_order") or []
    heading_case = publisher.get("lint", {}).get("heading_case") or publisher.get("docx", {}).get("numbering", {}).get("heading_case")
    enable_ko = detected_locale in {"ko", "mixed"}

    headings_seen = [name for name, _, _ in sections if name != "Preamble"]
    expected_pointer = 0
    normalized_expected = [normalize_heading(str(item)) for item in expected_order]
    for heading in headings_seen:
        normalized_heading = normalize_heading(heading)
        found_idx = None
        for idx in range(expected_pointer, len(normalized_expected)):
            if normalized_expected[idx] == normalized_heading:
                found_idx = idx
                break
        if found_idx is None:
            pos = full_text.find(f"# {heading}")
            expected_msg = expected_order[expected_pointer] if expected_pointer < len(expected_order) else "end of outline"
            add_finding(findings, full_text, offsets, max(pos, 0), "rule-08-structure-order", heading, f"Heading order deviates from publisher expectation: expected '{expected_msg}'", heading)
            break
        expected_pointer = found_idx + 1

    total_words = max(word_count(body), 1)
    for section_name, section_text, start_line in sections:
        if not section_text:
            continue
        section_pos = full_text.find(section_text)
        wc = max(word_count(section_text), 1)

        for pattern in ai_phrases:
            for match in pattern.finditer(section_text):
                add_finding(findings, full_text, offsets, section_pos + match.start(), "rule-01-ai-phrase", section_name, f"AI-tell phrase: '{match.group(0)}'", match.group(0))

        if enable_ko:
            for phrase in ko_rules.get("ai_tell_phrases", []):
                for match in re.finditer(re.escape(phrase), section_text):
                    add_finding(findings, full_text, offsets, section_pos + match.start(), "rule-21-ko-ai-phrase", section_name, f"한국어 AI 표현: '{match.group(0)}'", match.group(0))

        # A spaced hyphen between words reads as a dash; one opening a line is a list marker
        # (a key figure's basis line, a nested point), not a dash.
        # Table rows are data: a dash there marks an empty cell.
        prose = "\n".join(line for line in section_text.split("\n") if not line.lstrip().startswith("|"))
        if prose.count("—") >= 2 or re.search(r"\S - ", prose):
            add_finding(findings, full_text, offsets, max(section_pos, 0), "rule-02-em-dash-cluster", section_name, "Paragraph uses dash-heavy rhetorical style", section_text[:120])

        sentences = sentence_split(section_text)
        lengths = [word_count(sentence) for sentence in sentences]
        for sentence in sentences:
            hedge_hits = [word for word in re.findall(r"\b[a-zA-Z']+\b", sentence.lower()) if word in hedging_words]
            if len(hedge_hits) >= 2:
                pos = full_text.find(sentence)
                add_finding(findings, full_text, offsets, max(pos, 0), "rule-03-hedging-pileup", section_name, f"Sentence piles up hedging terms: {', '.join(hedge_hits)}", sentence[:120])

        if sentences:
            first_sentence = sentences[0].lower()
            last_sentence = sentences[-1].lower()
            if any(first_sentence.startswith(opener) for opener in stock_openers):
                pos = full_text.find(sentences[0])
                add_finding(findings, full_text, offsets, max(pos, 0), "rule-04-stock-boundary", section_name, "Section opens with a stock phrase", sentences[0][:120])
            if any(last_sentence.startswith(closer) for closer in stock_closers):
                pos = full_text.find(sentences[-1])
                add_finding(findings, full_text, offsets, max(pos, 0), "rule-04-stock-boundary", section_name, "Section closes with a stock phrase", sentences[-1][:120])
            if enable_ko:
                first_sentence_ko = sentences[0]
                last_sentence_ko = sentences[-1]
                if any(first_sentence_ko.startswith(opener) for opener in ko_rules.get("stock_openers", [])):
                    pos = full_text.find(sentences[0])
                    add_finding(findings, full_text, offsets, max(pos, 0), "rule-25-ko-stock-boundary", section_name, "한국어 섹션 도입 문구가 상투적입니다", sentences[0][:120])
                if any(last_sentence_ko.startswith(closer) for closer in ko_rules.get("stock_closers", [])):
                    pos = full_text.find(sentences[-1])
                    add_finding(findings, full_text, offsets, max(pos, 0), "rule-25-ko-stock-boundary", section_name, "한국어 섹션 마무리 문구가 상투적입니다", sentences[-1][:120])

        if enable_ko:
            eomi_hits = 0
            for marker in ko_rules.get("eomi_repetition", []):
                eomi_hits += section_text.count(marker)
            if eomi_hits >= 3:
                add_finding(findings, full_text, offsets, max(section_pos, 0), "rule-22-ko-eomi-repeat", section_name, f"종결어미 반복이 많습니다 ({eomi_hits} hits)", section_text[:120])

            for phrase in ko_rules.get("translationese", []):
                for match in re.finditer(re.escape(phrase), section_text):
                    add_finding(findings, full_text, offsets, section_pos + match.start(), "rule-23-ko-translationese", section_name, f"직역체 표현 감지: '{match.group(0)}'", match.group(0))

            for phrase in ko_rules.get("overused_loanwords", []):
                for match in re.finditer(re.escape(phrase), section_text):
                    add_finding(findings, full_text, offsets, section_pos + match.start(), "rule-24-ko-loanword", section_name, f"외래어 남용 감지: '{match.group(0)}'", match.group(0))

        # A reference list is a run of entries, not prose.
        if len(lengths) >= 4 and wc >= 80 and section_name.strip().lower() not in REFERENCE_SECTIONS:
            mean = sum(lengths) / len(lengths)
            variance = sum((item - mean) ** 2 for item in lengths) / len(lengths)
            if variance ** 0.5 < 6:
                add_finding(findings, full_text, offsets, max(section_pos, 0), "rule-05-sentence-variance", section_name, "Sentence lengths vary too little within this section", section_text[:120])

        ratio = passive_ratio(sentences)
        if section_name.lower() in {"methods", "results"} and ratio > 0.4:
            add_finding(findings, full_text, offsets, max(section_pos, 0), "rule-06-passive-ratio", section_name, f"Passive-voice ratio is high ({ratio:.0%})", section_text[:120])

        cites = citation_count(section_text)
        if section_name.lower() == "introduction" and cites == 0 and wc > 40:
            add_finding(findings, full_text, offsets, max(section_pos, 0), "rule-07-citation-density", section_name, "Introduction is under-cited for its length", section_text[:120])
        if section_name.lower() == "discussion" and cites == 0 and wc > 40:
            add_finding(findings, full_text, offsets, max(section_pos, 0), "rule-07-citation-density", section_name, "Discussion is under-cited for its length", section_text[:120])

        section_ratio = wc / total_words
        if total_words >= 250 and section_name.lower() == "introduction" and section_ratio > 0.20:
            add_finding(findings, full_text, offsets, max(section_pos, 0), "rule-09-section-balance", section_name, "Introduction is too large relative to the manuscript", section_text[:120])
        if total_words >= 250 and section_name.lower() == "methods" and section_ratio < 0.15:
            add_finding(findings, full_text, offsets, max(section_pos, 0), "rule-09-section-balance", section_name, "Methods section is too small relative to the manuscript", section_text[:120])

        # TTR falls with length, so it is read per subsection (the smallest headed unit), not over a whole chapter.
        parts = [part for part in re.split(r"(?m)^(?=#{2,6} )", section_text) if word_count(part) > 40]
        ttr = min((type_token_ratio(part) for part in parts), default=1.0)
        if ttr < 0.45:
            add_finding(findings, full_text, offsets, max(section_pos, 0), "rule-10-lexical-diversity", section_name, f"Lexical diversity is low (TTR={ttr:.2f})", section_text[:120])

        for pattern in adjective_stacks:
            for match in pattern.finditer(section_text):
                add_finding(findings, full_text, offsets, section_pos + match.start(), "rule-12-adjective-stack", section_name, "Promotional adjective stack detected", match.group(0))

        qualifier_hits = 0
        for token in re.findall(r"\b[a-zA-Z']+\b", section_text.lower()):
            if token in redundant:
                qualifier_hits += 1
        if qualifier_hits >= 3:
            add_finding(findings, full_text, offsets, max(section_pos, 0), "rule-13-redundant-qualifier", section_name, f"Redundant qualifiers appear {qualifier_hits} times", section_text[:120])

    for heading in headings_seen:
        if heading_case and not heading_case_ok(heading, heading_case):
            pos = full_text.find(f"# {heading}")
            add_finding(findings, full_text, offsets, max(pos, 0), "rule-11-heading-case", heading, f"Heading case does not match expected {heading_case}", heading)

    return findings, frontmatter, body, detected_locale


def apply_safe_fixes(body: str) -> str:
    body = fix_dashes(body)
    body = fix_quotes(body)
    body = fix_unit_spacing(body)
    body = fix_ellipsis(body)
    body = normalize_double_spaces(body)
    return body


def write_report(path: Path, findings: list[Finding], source: Path, publisher_name: str, detected_locale: str) -> None:
    lines = [
        "# Slop Lint Report",
        "",
        f"- Source: `{source}`",
        f"- Publisher: `{publisher_name}`",
        f"- Locale: `{detected_locale}`",
        f"- Findings: `{len(findings)}`",
        "",
    ]
    if not findings:
        lines.append("## Findings")
        lines.append("")
        lines.append("No findings.")
    else:
        lines.append("## Findings")
        lines.append("")
        for finding in findings:
            lines.append(f"- `{finding.rule_id}` L{finding.line}:C{finding.column} [{finding.section}] {finding.message} — {finding.excerpt}")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_submission_checklist(target_dir: Path, publisher_name: str, publisher: dict, frontmatter: dict) -> Path:
    target = target_dir / f"submission_checklist_{publisher_name}.md"
    required = publisher.get("frontmatter_required") or []
    lines = [
        f"# Submission checklist — {publisher_name}",
        "",
        "## Frontmatter",
    ]
    for item in required:
        present = "x" if frontmatter.get(item) else " "
        lines.append(f"- [{present}] {item}")
    lines.extend([
        "",
        "## Manual checks",
        "- [ ] ORCID / contributor metadata",
        "- [ ] Ethics / data-availability statements if needed",
        "- [ ] Figure and table counts checked against target journal",
    ])
    target.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return target


def audit_docx_cjk(path: Path) -> list[Finding]:
    findings: list[Finding] = []
    with zipfile.ZipFile(path) as zf:
        doc_xml = zf.read("word/document.xml").decode("utf-8")
    run_pattern = re.compile(r"<w:r(?:\s[^>]*)?>(.*?)</w:r>", re.DOTALL)
    text_pattern = re.compile(r"<w:t[^>]*>(.*?)</w:t>", re.DOTALL)
    eastasia_pattern = re.compile(r'w:eastAsia="([^"]+)"')
    for run in run_pattern.findall(doc_xml):
        texts = text_pattern.findall(run)
        if not texts:
            continue
        joined = "".join(texts)
        if not re.search(r"[가-힣]", joined):
            continue
        eastasia = eastasia_pattern.search(run)
        if eastasia is None or eastasia.group(1) in {"Batang", "Gulim", "MS Mincho"}:
            findings.append(Finding("rule-30-cjk-font-fallback", 0, 0, "DOCX", "Hangul run is missing explicit Pretendard-style eastAsia font pairing", joined[:80]))
    return findings


def audit_docx_design(path: Path, publisher: dict) -> list[Finding]:
    findings: list[Finding] = []
    banned_fonts = set((publisher.get("design") or {}).get("banned", {}).get("fonts") or [])
    with zipfile.ZipFile(path) as zf:
        styles_xml = zf.read("word/styles.xml").decode("utf-8")
        doc_xml = zf.read("word/document.xml").decode("utf-8")

    for font in banned_fonts:
        if font and font in styles_xml:
            findings.append(Finding("rule-32-banned-font", 0, 0, "DOCX", f"Banned font detected in styles.xml: {font}", font))

    margin_top = re.search(r'w:top="(\d+)"', doc_xml)
    expected_top_cm = ((publisher.get("docx") or {}).get("page") or {}).get("margin_top_cm")
    if margin_top and expected_top_cm is not None:
        actual = int(margin_top.group(1))
        expected = int(float(expected_top_cm) * 567)
        if abs(actual - expected) > 57:
            findings.append(Finding("rule-34-default-margin", 0, 0, "DOCX", f"Top margin differs from profile (got {actual}, expected ~{expected})", str(actual)))

    if "<w:tbl>" in doc_xml and 'w:insideV w:val="nil"' not in doc_xml:
        findings.append(Finding("rule-35-booktabs-table", 0, 0, "DOCX", "Booktabs insideV=nil marker not found", "insideV"))

    if 'w:highlight' in doc_xml:
        findings.append(Finding("rule-45-highlighted-text", 0, 0, "DOCX", "Highlighted text detected in document.xml", "highlight"))
    if re.search(r"<w:(ins|del|commentRangeStart)\b", doc_xml):
        findings.append(Finding("rule-46-tracked-changes", 0, 0, "DOCX", "Tracked changes or comments detected in document.xml", "tracked-changes"))
    if '...' in doc_xml:
        findings.append(Finding("rule-55-ellipsis", 0, 0, "DOCX", "ASCII ellipsis found in document.xml", "..."))
    return findings


# Family craft-floor subset for documents: OF-301 (text measure on the page width, advisory until a
# real .docx corpus calibrates it) and OF-302 (numeric table columns right-aligned).
DOCX_CRAFT = {
    "measure": {"latin_max": 90, "cjk_max": 38, "latin_em": 0.5, "cjk_em": 1.0, "cjk_share": 0.5, "default_pt": 11.0},
    "numeric": {"share": 0.7, "min_rows": 2},
}
_NUMERIC_STRIP = re.compile(r"[\s,.%▲▼+\-−±()$€£¥₩]")


def _numeric_cell(text: str) -> bool:
    stripped = _NUMERIC_STRIP.sub("", text)
    return bool(stripped) and stripped.isdigit()


def audit_docx_craft(path: Path) -> tuple[list[Finding], list[Finding]]:
    """Return (blocking, advisory) craft-floor findings for a generated DOCX."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    blocking: list[Finding] = []
    advisory: list[Finding] = []
    document = Document(str(path))

    t = DOCX_CRAFT["numeric"]
    for index, table in enumerate(document.tables, start=1):
        rows = table.rows[1:]
        if len(rows) < t["min_rows"]:
            continue
        for column in range(len(table.columns)):
            cells = [row.cells[column] for row in rows]
            filled = [cell for cell in cells if cell.text.strip()]
            numbers = [cell for cell in filled if _numeric_cell(cell.text)]
            if len(numbers) < t["min_rows"] or len(numbers) / max(1, len(filled)) < t["share"]:
                continue
            aligned = [[p.alignment for p in cell.paragraphs if p.text.strip()] for cell in numbers]
            explicit = sum(1 for values in aligned if any(value not in (None, WD_ALIGN_PARAGRAPH.RIGHT) for value in values))
            inherited = sum(1 for values in aligned if values and all(value is None for value in values))
            where = f"table {index} column {column + 1}"
            if explicit:
                blocking.append(Finding("rule-60-numeric-column-alignment", 0, 0, "DOCX", f"{where}: {explicit} number cell(s) not right-aligned (OF-302)", where))
            elif inherited:
                advisory.append(Finding("rule-60-numeric-column-alignment", 0, 0, "DOCX", f"{where}: alignment inherited from the style, not verified (OF-302)", where))

    m = DOCX_CRAFT["measure"]
    text = "".join(paragraph.text for paragraph in document.paragraphs)
    letters = [char for char in text if not char.isspace()]
    if letters and document.sections:
        cjk = sum(1 for char in letters if "\uac00" <= char <= "\ud7a3") / len(letters) >= m["cjk_share"]
        section = document.sections[0]
        usable_in = (section.page_width - section.left_margin - section.right_margin) / 914400
        with zipfile.ZipFile(path) as zf:
            styles = zf.read("word/styles.xml").decode("utf-8")
        normal = re.search(r'<w:style [^>]*w:styleId="Normal"[^>]*>(.*?)</w:style>', styles, re.S)
        size = re.search(r'<w:sz w:val="(\d+)"', normal.group(1) if normal else "") or re.search(r'<w:sz w:val="(\d+)"', styles)
        font_pt = int(size.group(1)) / 2 if size else m["default_pt"]
        per_line = usable_in * 72 / (font_pt * (m["cjk_em"] if cjk else m["latin_em"]))
        ceiling = m["cjk_max"] if cjk else m["latin_max"]
        if per_line > ceiling:
            advisory.append(Finding("rule-61-text-measure", 0, 0, "DOCX",
                                    f"about {per_line:.0f} {'Korean ' if cjk else ''}characters per line at {font_pt:g} pt over {usable_in:.2f} in; ceiling {ceiling} (OF-301)", f"{per_line:.0f}"))
    return blocking, advisory


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Non-destructive anti-slop lint for manuscript drafts", allow_abbrev=False)
    parser.add_argument("input", nargs="?", help="Input Markdown file")
    parser.add_argument("--publisher", required=True, help="Publisher profile from templates/registry.yaml")
    parser.add_argument("--registry", default=None, help="Path to templates/registry.yaml")
    parser.add_argument("--locale", choices=("auto", "en", "ko", "mixed"), default="auto")
    parser.add_argument("--report", required=True, help="Markdown report output path")
    parser.add_argument("--fix-whitespace", action="store_true")
    parser.add_argument("--audit-docx", default=None, help="Optional generated DOCX to audit for Korean CJK font fallback")
    parser.add_argument("--audit-output", default=None, help="Audit a generated DOCX artifact against design rules")
    parser.add_argument(
        "--kind",
        choices=("manuscript", "report"),
        default="manuscript",
        help="report skips the manuscript-only rules (citation density, section order, section balance) and the submission checklist",
    )
    return parser.parse_args()


# Rules that describe a journal manuscript's shape. A business report, proposal or plan
# has no Introduction/Methods outline to follow, so --kind report leaves them out.
REFERENCE_SECTIONS = {"references", "reference", "bibliography", "참고문헌", "참고 문헌"}
MANUSCRIPT_ONLY_RULES = {"rule-07-citation-density", "rule-08-structure-order", "rule-09-section-balance"}


def main() -> int:
    args = parse_args()
    script_dir = Path(__file__).resolve().parent
    report_path = Path(args.report).expanduser().resolve()
    registry = Path(args.registry).expanduser().resolve() if args.registry else script_dir.parent / "templates" / "registry.yaml"
    phrase_path = script_dir.parent / "references" / "slop_phrase_list.yaml"
    phrase_rules = load_phrase_rules(phrase_path)
    publisher = load_publisher(registry, args.publisher)

    if args.audit_output:
        audit_path = Path(args.audit_output).expanduser().resolve()
        if not audit_path.exists():
            print(f"Error: audit output not found: {audit_path}", file=sys.stderr)
            return 2
        findings = audit_docx_design(audit_path, publisher)
        if args.locale in {"ko", "mixed", "auto"}:
            findings.extend(audit_docx_cjk(audit_path))
        detected_locale = args.locale
        frontmatter = {}
        body = ""
        source = audit_path
    else:
        if not args.input:
            print("Error: input markdown path is required unless --audit-output is used", file=sys.stderr)
            return 2
        source = Path(args.input).expanduser().resolve()
        if not source.exists():
            print(f"Error: Input file not found: {source}", file=sys.stderr)
            return 2
        md_text = source.read_text(encoding="utf-8")
        findings, frontmatter, body, detected_locale = lint_text(md_text, publisher, phrase_rules, args.locale)

    if args.kind == "report":
        findings = [finding for finding in findings if finding.rule_id not in MANUSCRIPT_ONLY_RULES]

    if args.audit_docx and not args.audit_output:
        audit_path = Path(args.audit_docx).expanduser().resolve()
        if not audit_path.exists():
            print(f"Error: audit DOCX not found: {audit_path}", file=sys.stderr)
            return 2
        findings.extend(audit_docx_cjk(audit_path))

    if args.fix_whitespace and not args.audit_output:
        fixed = apply_safe_fixes(body)
        source.with_suffix(source.suffix + ".fixed.md").write_text(fixed, encoding="utf-8")

    write_report(report_path, findings, source, args.publisher, detected_locale)
    checklist_path = (
        write_submission_checklist(report_path.parent, args.publisher, publisher, frontmatter)
        if not args.audit_output and args.kind == "manuscript"
        else None
    )

    if findings:
        print(f"Findings: {len(findings)}")
        for finding in findings:
            print(f"{finding.rule_id} L{finding.line}:C{finding.column} [{finding.section}] {finding.message}")
        print(f"Report: {report_path}")
        if checklist_path:
            print(f"Checklist: {checklist_path}")
        return 1

    print("No findings.")
    print(f"Report: {report_path}")
    if checklist_path:
        print(f"Checklist: {checklist_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
