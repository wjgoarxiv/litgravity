#!/usr/bin/env python3
"""qa_docx.py — one pass/fail gate for a generated .docx and its Markdown source.

The document counterpart of lit-pptx's qa_deck.py. It combines:

  1. Integrity  — scripts/ooxml_integrity.py (the package opens,
                  relationships resolve, no NaN attribute values, python-docx reopens)
  2. Prose lint — slop_lint rules on the Markdown source (AI-tell phrases, hedging
                  pile-ups, stock openers/closers, adjective stacks, Korean rules);
                  --kind report drops the journal-manuscript outline rules
  3. Design     — slop_lint's DOCX audit (banned fonts, profile margins, booktabs
                  tables, highlights, tracked changes, ASCII ellipsis) and the Hangul
                  eastAsia font pairing check
  4. Content    — the document has body text, every Markdown heading reached it, no
                  unfilled blank ([제품명], [금액], XXX, TBD, ○○, {{…}}) is left, and
                  no YAML frontmatter leaked into the body
  5. Output     — docx_layout.py: headings, the title and the subtitle are noun-phrase
                  labels (heading.declarative), heading levels in order (heading.order),
                  components counted; with --layout the document is rendered through
                  soffice and its pages checked: fill.page, heading.stranded, table.split,
                  figure.split, component.variety and cover.block (FAILs on a tonality document);
                  the restraint checks of spec Amendments 4 always (color.accent-kinds, heading.ink,
                  heading.ratio, table.fill, component.budget, furniture.chip on a tonality document;
                  date.iso on every path); --compare OTHER.docx adds tonality.structure

Without --layout nothing is rendered; visual_audit.py writes the page images to look at.
A clean gate is defect-absence, not a verdict on how the pages read.

Usage:
    python3 qa_docx.py out.docx --source draft.md --layout [--kind report]        # tonality or plain
    python3 qa_docx.py out.docx --source draft.md --publisher korean-generic [--kind report]

Prints JSON; exit 0 on PASS, 1 on FAIL.
"""

from __future__ import annotations

import os as _os
import sys as _sys

_sys.dont_write_bytecode = True  # the installed skill directory stays read-only
_sys.path.insert(0, _os.path.join(_os.path.dirname(_os.path.abspath(__file__)), "..", "..", "lit-pptx", "scripts"))
from office_runtime_bootstrap import ensure_runtime  # noqa: E402

ensure_runtime(["docx", "markdown", "bs4", "yaml", "pymupdf"])

import argparse  # noqa: E402
import json  # noqa: E402
import re  # noqa: E402
import subprocess  # noqa: E402
import tempfile  # noqa: E402
from pathlib import Path  # noqa: E402

from docx import Document  # noqa: E402
from ooxml_integrity import check_package  # noqa: E402

SCRIPT_DIR = Path(__file__).resolve().parent
_sys.path.insert(0, str(SCRIPT_DIR))
import docx_layout  # noqa: E402
import slop_lint  # noqa: E402
from convert_md_to_docx import parse_frontmatter  # noqa: E402
from render_pages import to_pdf  # noqa: E402

HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
# A bracket with a digit in it is a citation marker ([1], [S3]); a blank has none.
PLACEHOLDER = re.compile(
    r"\[(?![^\]\n]*\d)[^\]\n]{1,24}\]|\bX{3,}\b|\bTBD\b|\bTODO\b|○○|OOO|\{\{[^}]*\}\}|_{4,}|＿{2,}"
)
FRONTMATTER_LEAK = re.compile(r"^(?:---|title:|author:|authors:|date:|subtitle:)\s*", re.M)


def source_headings(markdown_text: str) -> list[str]:
    headings, in_fence = [], False
    for line in markdown_text.splitlines():
        if line.lstrip().startswith(("```", "~~~")):
            in_fence = not in_fence
            continue
        match = None if in_fence else HEADING.match(line)
        if match:
            headings.append(re.sub(r"[*_`]", "", match.group(2)).strip())
    return headings


def check_content(docx_path: Path, markdown_text: str | None) -> dict:
    document = Document(str(docx_path))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    body_chars = sum(len(p) for p in paragraphs) + sum(
        len(cell.text) for table in document.tables for row in table.rows for cell in row.cells
    )
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    blanks = sorted({m.group(0) for text in paragraphs + cells for m in PLACEHOLDER.finditer(text)})
    leaked = [p for p in paragraphs[:8] if FRONTMATTER_LEAK.match(p)]
    missing = []
    if markdown_text:
        text = docx_layout.norm("".join(paragraphs + cells))  # a heading may stand in a side-by-side columns row
        for heading in source_headings(markdown_text):
            if docx_layout.norm(heading) not in text:
                missing.append(heading)
    return {
        "pass": body_chars > 0 and not missing and not blanks and not leaked,
        "placeholders": blanks[:20],
        "frontmatter_leak": leaked[:4],
        "paragraphs": len(paragraphs),
        "tables": len(document.tables),
        "body_chars": body_chars,
        "missing_headings": missing[:20],
    }


def lint(source: Path | None, docx_path: Path, publisher_name: str, kind: str, explicit_publisher: bool = True) -> dict:
    registry = SCRIPT_DIR.parent / "templates" / "registry.yaml"
    phrase_rules = slop_lint.load_phrase_rules(SCRIPT_DIR.parent / "references" / "slop_phrase_list.yaml")
    publisher = slop_lint.load_publisher(registry, publisher_name)
    findings = []
    locale = "auto"
    if source is not None:
        found, _front, _body, locale = slop_lint.lint_text(source.read_text(encoding="utf-8"), publisher, phrase_rules, "auto")
        findings.extend(found)
        if kind == "report":
            findings = [f for f in findings if f.rule_id not in slop_lint.MANUSCRIPT_ONLY_RULES]
        elif not explicit_publisher:
            # A tonality or plain manuscript has no publisher outline to follow.
            findings = [f for f in findings if f.rule_id != "rule-08-structure-order"]
    design = slop_lint.audit_docx_design(docx_path, publisher)
    if locale in {"ko", "mixed"} or source is None:
        design.extend(slop_lint.audit_docx_cjk(docx_path))
    craft_blocking, craft_advisory = slop_lint.audit_docx_craft(docx_path)
    design.extend(craft_blocking)
    as_dict = lambda f: {"rule": f.rule_id, "line": f.line, "section": f.section, "message": f.message, "excerpt": f.excerpt}  # noqa: E731
    return {
        "pass": not findings and not design,
        "locale": locale,
        "prose": [as_dict(f) for f in findings][:40],
        "design": [as_dict(f) for f in design][:40],
        "advisories": [as_dict(f) for f in craft_advisory][:40],
    }


def output_checks(docx_path: Path, source: Path | None, args) -> dict:
    """docx_layout's checks: structural always, the page checks when --layout renders the document."""
    front = parse_frontmatter(source.read_text(encoding="utf-8"))[0] if source else {}
    tonality = bool(args.tonality or front.get("tonality"))
    info = docx_layout.read_docx(docx_path)
    found = docx_layout.structural(info, front, args.kind == "manuscript", tonality)
    found += docx_layout.restraint(docx_path, tonality)
    result = {"components": info["components"], "rendered": False}
    if args.compare:
        mine, other = docx_layout.structure(docx_path), docx_layout.structure(args.compare.expanduser().resolve())
        differ = docx_layout.structure_diff(mine, other)
        result["structure"] = {"this": mine, "other": other, "differ": differ}
        if len(differ) < 3:
            found.append(docx_layout.finding("tonality.structure", "FAIL", None,
                                             f"the two documents differ in {len(differ)} structural feature(s): {', '.join(differ) or 'none'} (at least 3)",
                                             "Two tonalities differ in structure, not colour: title block, summary form, numbering, component set, running head, contents."))
    if args.layout:
        with tempfile.TemporaryDirectory(prefix="lit-docx-layout-") as tmp:
            try:
                pdf, how = to_pdf(docx_path, Path(tmp))
            except (RuntimeError, subprocess.TimeoutExpired) as error:
                pdf, how = None, f"render failed: {error}"
            if pdf is None:
                result["render"] = how  # the reply must say the page checks did not run
            else:
                page_found, pages = docx_layout.paged(info, docx_layout.read_pdf(pdf), tonality, front)
                found += page_found
                result.update(pages, rendered=True)
    result["findings"] = [f for f in found if f["severity"] == "FAIL"]
    result["advisories"] = [f for f in found if f["severity"] != "FAIL"]
    return result


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Pass/fail gate for a generated DOCX.")
    parser.add_argument("docx", type=Path)
    parser.add_argument("--source", type=Path, default=None, help="Markdown the DOCX was built from")
    parser.add_argument(
        "--publisher",
        default=None,
        help="Profile the DOCX was built with (korean-generic, elsevier, acs, ieee, nature). Omit for the plain profile.",
    )
    parser.add_argument("--kind", choices=("report", "manuscript"), default="report")
    parser.add_argument("--layout", action="store_true", help="render through soffice and check the pages (fill, stranded headings, splits, component variety)")
    parser.add_argument("--tonality", default=None, help="the tonality the DOCX was built with, when the source frontmatter does not name it")
    parser.add_argument("--compare", type=Path, default=None, help="the same source built in another tonality: fail when the two differ in fewer than three structural features")
    args = parser.parse_args(argv)

    docx_path = args.docx.expanduser().resolve()
    source = args.source.expanduser().resolve() if args.source else None
    integrity = check_package(docx_path)
    content = check_content(docx_path, source.read_text(encoding="utf-8") if source else None) if integrity["pass"] else {"pass": False}
    # The plain profile has no registry entry; its prose is still linted with the
    # korean-generic phrase rules, and the design audit is limited to what every
    # profile shares.
    lint_report = lint(source, docx_path, args.publisher or "korean-generic", args.kind, args.publisher is not None) if integrity["pass"] else {"pass": False}
    if args.publisher is None and lint_report.get("design"):
        shared = {"rule-45-highlighted-text", "rule-46-tracked-changes", "rule-55-ellipsis", "rule-60-numeric-column-alignment"}
        lint_report["design"] = [f for f in lint_report["design"] if f["rule"] in shared]
        lint_report["pass"] = not lint_report["prose"] and not lint_report["design"]

    output = output_checks(docx_path, source, args) if integrity["pass"] else {"findings": [], "advisories": []}

    reasons = []
    if not integrity["pass"]:
        reasons.append("integrity: " + "; ".join(integrity["problems"][:5]))
    if content.get("placeholders"):
        reasons.append(f"content: unfilled blanks {content['placeholders'][:5]} — fill them with realistic values and label them as examples or assumptions")
    if content.get("frontmatter_leak"):
        reasons.append("content: YAML frontmatter printed as body text — convert with the converter that reads frontmatter")
    if not content.get("pass") and not content.get("placeholders") and not content.get("frontmatter_leak"):
        reasons.append("content: empty document or headings missing from the output")
    if lint_report.get("prose"):
        reasons.append(f"prose lint: {len(lint_report['prose'])} finding(s)")
    if lint_report.get("design"):
        reasons.append(f"design audit: {len(lint_report['design'])} finding(s)")
    if output["findings"]:
        reasons.append(f"output checks: {', '.join(sorted({f['check'] for f in output['findings']}))}")

    report = {
        "file": str(docx_path),
        "pass": not reasons,
        "failure_reasons": reasons,
        "profile": args.publisher or "plain",
        "kind": args.kind,
        "integrity": {"pass": integrity["pass"], "problems": integrity["problems"][:20]},
        "content": content,
        "lint": lint_report,
        "output": output,
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["pass"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
