#!/usr/bin/env python3
"""Honest, renderer-agnostic QA gate for generated PPTX decks.

Brand-neutral: works for any enrolled template. Combines these checks and
exits non-zero if ANY real defect is found:

  1. Layout issues  — layout_inventory.py (frame overflow / slide overflow / off-slide)
  2. Anti-slop      — validate_pptx.py (forbidden/placeholder terms, default fonts)
  3. WCAG contrast  — text colour vs its fill (body < 4.5:1, large < 3:1)
  4. Integrity      — scripts/ooxml_integrity.py (package opens,
                      relationships resolve, no NaN attribute values)
  5. Deck craft     — deck_craft.py (sparse slides, table-only decks, shapes off
                      the canvas, stray empty boxes, [placeholder] blanks, long
                      foreign-language source lines on a Korean deck)
  6. Craft floor    — craft_extras.py (OF-101..OF-109 per slide; OF-110..OF-113
                      over the deck: title-treatment count, composition variety,
                      empty bottom band, one frame per title treatment; OF-114
                      titles and the cover subtitle as labels, not sentences;
                      OF-115 no empty region; OF-117 no light figure card on a
                      dark ground; OF-116 with --sibling: another tonality of
                      the same source draws another skeleton)

No LibreOffice / rendering engine required: everything is read from the OOXML
via python-pptx. (A renderer is only ever needed for the *visual* pass, which
lives in scripts/render_pages.py and needs soffice.)
"""

from __future__ import annotations

import os as _os
import sys as _sys

_sys.dont_write_bytecode = True  # the installed skill directory stays read-only
_sys.path.insert(0, _os.path.dirname(_os.path.abspath(__file__)))
from office_runtime_bootstrap import ensure_runtime  # noqa: E402

ensure_runtime(["pptx"])

import argparse
import json
import re
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
VALIDATE = ROOT / "scripts" / "validate_pptx.py"
INVENTORY = ROOT / "scripts" / "layout_inventory.py"

# WCAG 2.x contrast floors.
BODY_MIN = 4.5
LARGE_MIN = 3.0


# ── WCAG contrast ──────────────────────────────────────────────────────────
def _rel_luminance(rgb: tuple[int, int, int]) -> float:
    def chan(c: int) -> float:
        s = c / 255.0
        return s / 12.92 if s <= 0.03928 else ((s + 0.055) / 1.055) ** 2.4

    r, g, b = (chan(x) for x in rgb)
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def _ratio(fg: tuple[int, int, int], bg: tuple[int, int, int]) -> float:
    l1, l2 = _rel_luminance(fg), _rel_luminance(bg)
    hi, lo = max(l1, l2), min(l1, l2)
    return (hi + 0.05) / (lo + 0.05)


def _hex_to_rgb(value: Any) -> tuple[int, int, int] | None:
    if isinstance(value, RGBColor):
        value = str(value)
    if not isinstance(value, str):
        return None
    v = value.lstrip("#")
    if len(v) != 6:
        return None
    try:
        return (int(v[0:2], 16), int(v[2:4], 16), int(v[4:6], 16))
    except ValueError:
        return None


def _shape_fill_rgb(shape) -> tuple[int, int, int] | None:
    fill = getattr(shape, "fill", None)
    if fill is None or getattr(fill, "type", None) is None:
        return None
    try:
        rgb = fill.fore_color.rgb
    except (TypeError, AttributeError):
        return None
    return _hex_to_rgb(rgb)


def _is_large(size_pt: float | None, bold: bool) -> bool:
    if size_pt is None:
        return False
    return size_pt >= 18 or (bold and size_pt >= 14)


def _bbox(shape):
    try:
        l, t, w, h = int(shape.left), int(shape.top), int(shape.width), int(shape.height)
        return (l, t, l + w, t + h)
    except (TypeError, ValueError):
        return None


def _slide_ground(slide) -> tuple[int, int, int]:
    """The slide's own solid background, white when it has none (or inherits one)."""
    try:
        fill = slide.background.fill
        if fill.type == 1:  # solid
            return _hex_to_rgb(fill.fore_color.rgb) or (255, 255, 255)
    except (AttributeError, TypeError, ValueError):
        pass
    return (255, 255, 255)


def _resolve_bg(shape, idx, filled, ground=(255, 255, 255)):
    """Effective background for a text shape: its own fill, else the topmost
    filled shape drawn *behind* it (lower z-order) that contains its centre,
    else the slide background. Prevents false positives like white text over a
    charcoal band, or light text on a dark slide."""
    own = _shape_fill_rgb(shape)
    if own is not None:
        return own
    box = _bbox(shape)
    if box is not None:
        cx, cy = (box[0] + box[2]) / 2, (box[1] + box[3]) / 2
        best = None
        for f_idx, f_rgb, f_box in filled:
            if f_idx >= idx or f_box is None:
                continue
            if f_box[0] <= cx <= f_box[2] and f_box[1] <= cy <= f_box[3]:
                best = f_rgb  # later (higher idx) wins -> topmost behind text
        if best is not None:
            return best
    return ground


def check_contrast(path: Path) -> dict[str, Any]:
    """Flag explicit-colour text that falls under the WCAG floor on its fill."""
    prs = Presentation(str(path))
    violations: list[dict[str, Any]] = []
    checked = 0
    for s_idx, slide in enumerate(prs.slides, start=1):
        shapes = list(slide.shapes)
        filled = [(i, rgb, _bbox(sh)) for i, sh in enumerate(shapes) if (rgb := _shape_fill_rgb(sh)) is not None]
        ground = _slide_ground(slide)
        for idx, shape in enumerate(shapes):
            if not getattr(shape, "has_text_frame", False):
                continue
            bg = _resolve_bg(shape, idx, filled, ground)
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if not run.text.strip():
                        continue
                    try:
                        fg = _hex_to_rgb(run.font.color.rgb)
                    except (TypeError, AttributeError):
                        fg = None
                    if fg is None:
                        continue  # theme/inherited colour — cannot judge, skip
                    checked += 1
                    size = float(run.font.size.pt) if run.font.size is not None else None
                    bold = bool(run.font.bold)
                    ratio = _ratio(fg, bg)
                    floor = LARGE_MIN if _is_large(size, bold) else BODY_MIN
                    if ratio < floor:
                        violations.append({
                            "slide": s_idx,
                            "text": run.text.strip()[:60],
                            "fg": "#%02X%02X%02X" % fg,
                            "bg": "#%02X%02X%02X" % bg,
                            "ratio": round(ratio, 2),
                            "floor": floor,
                        })
    return {"pass": not violations, "checked_runs": checked, "violations": violations[:30]}


# ── picture / table overlap ─────────────────────────────────────────────────
# The layout inventory's overlap list is plain box geometry with no z-order, so it
# cannot tell a chart drawn over a data table from text sitting on its card. This check covers
# PICTURE and TABLE (graphicFrame) shapes against every other content-bearing
# shape and fails on any real overlap.
OVERLAP_TOL_IN = 0.06  # inches; ignore hairline touches / rounding
MIN_TEXT_PT = 6.0


def _shape_kind(shape) -> str | None:
    """Classify a shape for overlap purposes: 'picture', 'table', 'text', or
    None for decorative/empty shapes (backgrounds, rules, blank frames)."""
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return "picture"
    if getattr(shape, "has_table", False):
        return "table"
    if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
        return "text"
    return None


def _bbox_in(shape) -> tuple[float, float, float, float] | None:
    try:
        l = shape.left / 914400.0
        t = shape.top / 914400.0
        w = shape.width / 914400.0
        h = shape.height / 914400.0
    except (TypeError, ValueError):
        return None
    return (l, t, l + w, t + h)


def _is_full_bleed(box, page_area: float) -> bool:
    """A shape covering essentially the whole canvas is a ground, not an occluder."""
    return ((box[2] - box[0]) * (box[3] - box[1])) >= 0.92 * page_area


def _overlap_extent(box_a, box_b) -> tuple[float, float]:
    return (min(box_a[2], box_b[2]) - max(box_a[0], box_b[0]),
            min(box_a[3], box_b[3]) - max(box_a[1], box_b[1]))


def check_semantic_overlaps(path: Path) -> dict[str, Any]:
    """Text running over other text — the one occlusion nothing else catches.

    Free positioning lets an author drop a box anywhere, so two text frames can
    now land on top of each other and produce a slide where neither is readable.
    The layout inventory lists box overlaps but cannot tell a card's own text
    from two frames colliding, so its count stays informational.

    Every other pairing already has a more careful owner, and duplicating them
    here would only reintroduce their false positives:

      picture or table drawn over content   -> check_overlaps (z-order aware)
      declared figure or table occluded     -> check_figure_table_occlusion
      text on a card, panel or backdrop     -> check_contrast, on its real fill

    That last one is why an undeclared picture is not treated as content. A
    gradient circle behind a cover title is a ground exactly as a navy panel is;
    a picture only carries content once its description declares it as evidence,
    and check_figure_table_occlusion is what reads that declaration.
    """
    prs = Presentation(str(path))
    page_area = (prs.slide_width / 914400.0) * (prs.slide_height / 914400.0)
    violations: list[dict[str, Any]] = []

    for s_idx, slide in enumerate(prs.slides, start=1):
        texts = [
            (_bbox_in(sh), sh) for sh in slide.shapes
            if _shape_kind(sh) == "text" and _bbox_in(sh) is not None
        ]
        for i in range(len(texts)):
            for j in range(i + 1, len(texts)):
                box_a, sh_a = texts[i]
                box_b, sh_b = texts[j]
                if _is_full_bleed(box_a, page_area) or _is_full_bleed(box_b, page_area):
                    continue
                ow, oh = _overlap_extent(box_a, box_b)
                if ow <= OVERLAP_TOL_IN or oh <= OVERLAP_TOL_IN:
                    continue
                violations.append({
                    "slide": s_idx,
                    "a": "text",
                    "b": "text",
                    "area": round(ow * oh, 2),
                    "a_text": _shape_text(sh_a)[:40],
                    "b_text": _shape_text(sh_b)[:40],
                })

    return {"pass": not violations, "violations": violations[:30]}


def check_overlaps(path: Path) -> dict[str, Any]:
    """A PICTURE or TABLE must never be drawn on top of (higher z-order than)
    another content-bearing shape. Decorations are drawn first (behind), so they
    are never the 'top' shape and never false-flag."""
    prs = Presentation(str(path))
    page_area = (prs.slide_width / 914400.0) * (prs.slide_height / 914400.0)
    violations: list[dict[str, Any]] = []
    for s_idx, slide in enumerate(prs.slides, start=1):
        shapes = []
        for z, sh in enumerate(slide.shapes):
            box = _bbox_in(sh)
            if box is None:
                continue
            shapes.append((z, sh, _shape_kind(sh), box))

        def is_full_bleed(b):
            return _is_full_bleed(b, page_area)
        for z, sh, kind, box in shapes:
            if kind not in ("picture", "table"):
                continue
            for z2, sh2, kind2, box2 in shapes:
                if z2 >= z or kind2 is None:
                    continue  # only shapes drawn *behind* this picture/table
                # A full-bleed background image (mesh/gradient wash) is never an
                # occlusion victim — anything may sit on top of it by design.
                if kind2 == "picture" and is_full_bleed(box2):
                    continue
                # Picture-over-picture is intentional decorative layering.
                if kind == "picture" and kind2 == "picture":
                    continue
                ow = min(box[2], box2[2]) - max(box[0], box2[0])
                oh = min(box[3], box2[3]) - max(box[1], box2[1])
                if ow > OVERLAP_TOL_IN and oh > OVERLAP_TOL_IN:
                    violations.append({
                        "slide": s_idx,
                        "top": kind,
                        "under": kind2,
                        "area": round(ow * oh, 2),
                    })
    return {"pass": not violations, "violations": violations[:30]}


# ── objective executive/research checks ────────────────────────────────────
def _intersection(box_a, box_b) -> tuple[float, float]:
    return (
        min(box_a[2], box_b[2]) - max(box_a[0], box_b[0]),
        min(box_a[3], box_b[3]) - max(box_a[1], box_b[1]),
    )


def _iter_shape_runs(shape):
    if getattr(shape, "has_text_frame", False):
        for paragraph in shape.text_frame.paragraphs:
            yield from paragraph.runs
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            for cell in row.cells:
                for paragraph in cell.text_frame.paragraphs:
                    yield from paragraph.runs


def _shape_text(shape) -> str:
    return "\n".join(run.text for run in _iter_shape_runs(shape) if run.text).strip()


def _semantic_label(shape) -> str:
    values = [getattr(shape, "name", "") or ""]
    try:
        c_nv_pr = shape._element.xpath(".//p:cNvPr")[0]
        values.extend([c_nv_pr.get("name", ""), c_nv_pr.get("descr", ""), c_nv_pr.get("title", "")])
    except (AttributeError, IndexError):
        pass
    return " ".join(value for value in values if value).strip()


def check_side_stripes(path: Path) -> dict[str, Any]:
    """Detect a narrow filled bar attached to the vertical edge of a larger card."""
    prs = Presentation(str(path))
    violations: list[dict[str, Any]] = []
    for slide_idx, slide in enumerate(prs.slides, start=1):
        filled = [
            (idx, shape, _bbox_in(shape), _shape_fill_rgb(shape))
            for idx, shape in enumerate(slide.shapes)
            if _bbox_in(shape) is not None and _shape_fill_rgb(shape) is not None
        ]
        for stripe_idx, stripe, stripe_box, stripe_rgb in filled:
            sw = stripe_box[2] - stripe_box[0]
            sh = stripe_box[3] - stripe_box[1]
            if sw > 0.15 or sh < 0.5 or sh / max(sw, 0.001) < 5:
                continue
            for card_idx, card, card_box, card_rgb in filled:
                if card_idx == stripe_idx or card_rgb == stripe_rgb:
                    continue
                cw = card_box[2] - card_box[0]
                ch = card_box[3] - card_box[1]
                if cw < 1.0 or ch < 0.5 or sh < ch * 0.7:
                    continue
                same_vertical_band = abs(stripe_box[1] - card_box[1]) <= 0.08 and abs(stripe_box[3] - card_box[3]) <= 0.08
                on_edge = abs(stripe_box[0] - card_box[0]) <= 0.08 or abs(stripe_box[2] - card_box[2]) <= 0.08
                if same_vertical_band and on_edge:
                    violations.append({
                        "slide": slide_idx,
                        "stripe_shape": stripe_idx + 1,
                        "card_shape": card_idx + 1,
                        "stripe_width_in": round(sw, 3),
                    })
                    break
    return {"pass": not violations, "violations": violations[:30]}


def check_text_size_floor(path: Path) -> dict[str, Any]:
    """Fail only on explicit text below the absolute 6pt export floor."""
    prs = Presentation(str(path))
    violations: list[dict[str, Any]] = []
    checked = 0
    for slide_idx, slide in enumerate(prs.slides, start=1):
        for shape_idx, shape in enumerate(slide.shapes, start=1):
            for run in _iter_shape_runs(shape):
                if not run.text.strip() or run.font.size is None:
                    continue
                checked += 1
                size = float(run.font.size.pt)
                if size < MIN_TEXT_PT:
                    violations.append({
                        "slide": slide_idx,
                        "shape": shape_idx,
                        "text": run.text.strip()[:60],
                        "size_pt": round(size, 2),
                        "floor_pt": MIN_TEXT_PT,
                    })
    return {"pass": not violations, "checked_runs": checked, "violations": violations[:30]}


def check_figure_table_occlusion(path: Path) -> dict[str, Any]:
    """Flag content shapes drawn over a non-background figure or table."""
    prs = Presentation(str(path))
    violations: list[dict[str, Any]] = []
    for slide_idx, slide in enumerate(prs.slides, start=1):
        page_area = (prs.slide_width / 914400.0) * (prs.slide_height / 914400.0)
        shapes = [
            (idx, shape, _shape_kind(shape), _bbox_in(shape))
            for idx, shape in enumerate(slide.shapes)
            if _bbox_in(shape) is not None
        ]
        for victim_idx, victim, victim_kind, victim_box in shapes:
            if victim_kind not in ("picture", "table"):
                continue
            # A generic picture may intentionally be a background or decorated
            # asset. Only pictures explicitly declared as evidence/figures have
            # objective occlusion semantics; tables are always data objects.
            if victim_kind == "picture" and not _EVIDENCE_MARKER.search(_semantic_label(victim)):
                continue
            victim_area = (victim_box[2] - victim_box[0]) * (victim_box[3] - victim_box[1])
            if victim_kind == "picture" and victim_area >= 0.92 * page_area:
                continue
            for top_idx, top, top_kind, top_box in shapes:
                if top_idx <= victim_idx or top_kind is None:
                    continue
                if victim_kind == "picture" and top_kind == "picture":
                    continue
                ow, oh = _intersection(victim_box, top_box)
                if ow <= OVERLAP_TOL_IN or oh <= OVERLAP_TOL_IN:
                    continue
                overlap_area = ow * oh
                if overlap_area < max(0.08, victim_area * 0.03):
                    continue
                violations.append({
                    "slide": slide_idx,
                    "victim": victim_kind,
                    "victim_shape": victim_idx + 1,
                    "occluding": top_kind,
                    "occluding_shape": top_idx + 1,
                    "area": round(overlap_area, 2),
                })
    return {"pass": not violations, "violations": violations[:30]}


_EVIDENCE_MARKER = re.compile(
    r"(?:^|\b)(?:figure|evidence|source[ -]?capture|그림|도표|근거|출처\s*캡처)\s*(?::|\d|\b)",
    re.IGNORECASE,
)
_CAPTION_MARKER = re.compile(r"(?:그림|도|figure|fig\.)\s*\d+", re.IGNORECASE)
_SOURCE_MARKER = re.compile(r"(?:출처|source)\s*:", re.IGNORECASE)


def check_evidence_binding(path: Path) -> dict[str, Any]:
    """Require visible, adjacent caption and source text for declared evidence figures."""
    prs = Presentation(str(path))
    violations: list[dict[str, Any]] = []
    checked = 0
    for slide_idx, slide in enumerate(prs.slides, start=1):
        shapes = list(slide.shapes)
        for shape_idx, shape in enumerate(shapes, start=1):
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE or not _EVIDENCE_MARKER.search(_semantic_label(shape)):
                continue
            checked += 1
            figure_box = _bbox_in(shape)
            nearby: list[str] = []
            page_area = (prs.slide_width / 914400.0) * (prs.slide_height / 914400.0)
            if figure_box is not None and _is_full_bleed(figure_box, page_area):
                # A full-bleed picture carries its caption on the panel laid over it.
                nearby = [text for other in shapes if other is not shape and (text := _shape_text(other))]
            elif figure_box is not None:
                figure_width = figure_box[2] - figure_box[0]
                for other in shapes:
                    text = _shape_text(other)
                    other_box = _bbox_in(other)
                    if not text or other_box is None:
                        continue
                    horizontal = min(figure_box[2], other_box[2]) - max(figure_box[0], other_box[0])
                    vertical_gap = other_box[1] - figure_box[3]
                    if horizontal >= figure_width * 0.2 and -0.04 <= vertical_gap <= 1.1:
                        nearby.append(text)
            bound_text = "\n".join(nearby)
            missing = []
            if not _CAPTION_MARKER.search(bound_text):
                missing.append("caption")
            if not _SOURCE_MARKER.search(bound_text):
                missing.append("source")
            if missing:
                violations.append({
                    "slide": slide_idx,
                    "shape": shape_idx,
                    "label": _semantic_label(shape)[:100],
                    "missing": missing,
                })
    return {"pass": not violations, "checked_figures": checked, "violations": violations[:30]}


_APPENDIX_MARKER = re.compile(r"\bappendix\b|부록", re.IGNORECASE)
_RAW_URL = re.compile(r"https?://[^\s<>\]\[)}]+", re.IGNORECASE)
_DOI_VALUE = re.compile(r"(?<!doi\.org/)\b10\.\d{4,9}/[^\s<>\]\[)}]+", re.IGNORECASE)
_LINK_FIELD = re.compile(
    r"\bdoi\b|canonical\s+link|source\s+link|정식\s*링크|원문(?:\s*링크)?",
    re.IGNORECASE,
)
_SOURCE_LOCATOR = re.compile(r"\bsource\b|출처", re.IGNORECASE)


def _appendix_link_units(shape):
    """Yield visible text and run-level hyperlink targets per paragraph/cell."""
    if getattr(shape, "has_table", False):
        headers = [cell.text.strip() for cell in shape.table.rows[0].cells]
        for row_idx, row in enumerate(shape.table.rows):
            if row_idx == 0:
                continue
            for column_idx, cell in enumerate(row.cells):
                for paragraph in cell.text_frame.paragraphs:
                    runs = [(run.text, run.hyperlink.address) for run in paragraph.runs if run.text]
                    text = "".join(run_text for run_text, _ in runs).strip()
                    if text:
                        yield (
                            text,
                            [address for _, address in runs if address],
                            bool(
                                _LINK_FIELD.search(headers[column_idx])
                                or (
                                    _SOURCE_LOCATOR.search(headers[column_idx])
                                    and any(address for _, address in runs)
                                )
                            ),
                        )
        return
    if getattr(shape, "has_text_frame", False):
        for paragraph in shape.text_frame.paragraphs:
            runs = [(run.text, run.hyperlink.address) for run in paragraph.runs if run.text]
            text = "".join(run_text for run_text, _ in runs).strip()
            if text:
                run_links = [address for _, address in runs if address]
                yield text, run_links, bool(
                    _LINK_FIELD.search(text)
                    or (_SOURCE_LOCATOR.search(text) and run_links)
                )


def _slide_hyperlink_targets(slide) -> list[str]:
    targets = []
    for relationship in slide.part.rels.values():
        if relationship.reltype == RT.HYPERLINK and relationship.is_external:
            targets.append(relationship.target_ref)
    return targets


def check_appendix_links(path: Path) -> dict[str, Any]:
    """Appendix DOI, canonical-link, and raw-URL fields must be clickable."""
    prs = Presentation(str(path))
    violations: list[dict[str, Any]] = []
    declared = 0
    relationship_targets: list[str] = []
    for slide_idx, slide in enumerate(prs.slides, start=1):
        slide_text = "\n".join(_shape_text(shape) for shape in slide.shapes)
        if not _APPENDIX_MARKER.search(slide_text):
            continue
        slide_targets = _slide_hyperlink_targets(slide)
        relationship_targets.extend(slide_targets)
        for shape_idx, shape in enumerate(slide.shapes, start=1):
            for text, run_links, is_link_field in _appendix_link_units(shape):
                raw_urls = [url.rstrip(".,;:") for url in _RAW_URL.findall(text)]
                doi_values = [] if raw_urls else [doi.rstrip(".,;:") for doi in _DOI_VALUE.findall(text)]
                candidates = raw_urls or doi_values
                if not candidates and is_link_field:
                    candidates = [run_links[0]] if run_links else [None]
                for candidate in candidates:
                    declared += 1
                    if candidate is None:
                        linked = False
                    elif candidate.startswith(("http://", "https://")):
                        linked = candidate in run_links and candidate in slide_targets
                    else:
                        linked = any(
                            candidate.lower() in target.lower() and target in run_links
                            for target in slide_targets
                        )
                    if not linked:
                        violations.append({
                            "slide": slide_idx,
                            "shape": shape_idx,
                            "text": text[:100],
                            "locator": candidate,
                            "missing": "OOXML hyperlink relationship",
                        })
    return {
        "pass": not violations,
        "declared_links": declared,
        "hyperlink_relationships": len(relationship_targets),
        "relationship_targets": relationship_targets[:20],
        "violations": violations[:30],
    }


def check_objective_quality(path: Path) -> dict[str, Any]:
    checks = {
        "side_stripes": check_side_stripes(path),
        "text_size_floor": check_text_size_floor(path),
        "figure_table_occlusion": check_figure_table_occlusion(path),
        "evidence_binding": check_evidence_binding(path),
        "appendix_links": check_appendix_links(path),
    }
    return {"pass": all(check["pass"] for check in checks.values()), **checks}


# ── inventory issues ───────────────────────────────────────────────────────
def load_inventory_issues(pptx: Path) -> dict[str, Any]:
    with tempfile.NamedTemporaryFile(suffix=".json", delete=False) as tmp:
        tmp_path = Path(tmp.name)
    try:
        subprocess.run(
            [sys.executable, str(INVENTORY), str(pptx), str(tmp_path), "--issues-only"],
            capture_output=True, text=True, cwd=ROOT, timeout=120,
        )
        return json.loads(tmp_path.read_text(encoding="utf-8")) if tmp_path.exists() and tmp_path.stat().st_size else {}
    finally:
        tmp_path.unlink(missing_ok=True)


def summarize_inventory(issues: dict[str, Any]) -> dict[str, Any]:
    if not issues:
        # No inventory is not a clean inventory: report it so the gate fails.
        return {"available": False, "issue_shapes": 0, "overflow_shapes": 0, "slide_overflow_shapes": 0,
                "frame_overflow_shapes": 0, "off_slide_shapes": 0, "overlap_shapes": 0, "details": []}
    shape_count = overflow = slide_overflow = frame_overflow = off_slide = overlap = 0
    details: list[dict[str, Any]] = []
    for slide in issues.get("slides", []):
        for shape in slide.get("shapes", []):
            found = shape.get("issues") or {}
            if not found:
                continue
            shape_count += 1
            defect = False
            if "frame_overflow" in found:
                frame_overflow += 1
                defect = True
            if "slide_overflow" in found:
                slide_overflow += 1
                defect = True
            if "off_slide" in found:
                off_slide += 1
                defect = True
            if defect:
                overflow += 1
                details.append({"slide": slide.get("slide"), "shape": shape.get("name"),
                                **{k: v for k, v in found.items() if k != "overlaps"}})
            # Reported only — see qa(): raw box overlap has no notion of z-order.
            if found.get("overlaps"):
                overlap += 1
    return {
        "available": True,
        "issue_shapes": shape_count,
        "overflow_shapes": overflow,
        "slide_overflow_shapes": slide_overflow,
        "frame_overflow_shapes": frame_overflow,
        "off_slide_shapes": off_slide,
        "overlap_shapes": overlap,
        "details": details[:30],
    }


def run_validate(pptx: Path) -> dict[str, Any]:
    result = subprocess.run(
        [sys.executable, str(VALIDATE), str(pptx)],
        capture_output=True, text=True, cwd=ROOT, timeout=120,
    )
    try:
        return json.loads(result.stdout)
    except json.JSONDecodeError:
        return {"pass": False, "error": "validate_pptx produced no JSON", "stderr": result.stderr[:400]}


# ── gate ───────────────────────────────────────────────────────────────────
def qa(pptx: Path, sibling: Path | None = None) -> dict[str, Any]:
    from ooxml_integrity import check_package
    from deck_craft import analyse as analyse_craft
    from craft_extras import analyse as analyse_extras, compare_skeletons

    integrity = check_package(pptx)
    craft = analyse_craft(pptx)
    extras = analyse_extras(pptx)
    # Another tonality of the same source, when given: the two must not draw one skeleton (OF-116).
    if sibling is not None:
        extras["findings"] += compare_skeletons(pptx, sibling)
    validate_report = run_validate(pptx)
    inv_summary = summarize_inventory(load_inventory_issues(pptx))
    contrast = check_contrast(pptx)
    overlaps = check_overlaps(pptx)
    semantic = check_semantic_overlaps(pptx)
    objective = check_objective_quality(pptx)

    reasons: list[str] = []
    if not integrity["pass"]:
        reasons.append("integrity: " + "; ".join(integrity["problems"][:5]))
    by_check: dict[str, list] = {}
    for finding in craft["findings"]:
        by_check.setdefault(finding["check"], []).append(finding)
    for check, found in by_check.items():
        slides = sorted({f["slide"] for f in found if f["slide"] is not None})
        where = f" on slide(s) {', '.join(map(str, slides))}" if slides else ""
        reasons.append(f"craft {check}{where} — {found[0]['hint']}")
    extras_high: dict[str, list] = {}
    for finding in extras["findings"]:
        if finding["severity"] == "HIGH":
            extras_high.setdefault(finding["check"], []).append(finding)
    for check, found in extras_high.items():
        slides = sorted({f["slide"] for f in found if f["slide"] is not None})
        where = f" on slide(s) {', '.join(map(str, slides))}" if slides else f" ({found[0]['detail']})"
        reasons.append(f"craft {check}{where} — {found[0]['hint']}")
    if not validate_report.get("pass", False):
        reasons.append("anti-slop: " + ", ".join(
            k for k, v in validate_report.get("checks", {}).items() if isinstance(v, dict) and not v.get("pass", True)
        ) or "anti-slop: validate_pptx failed")
    # Fail only on real layout defects: text spilling out of its frame, text or a
    # grown table running past the canvas, a text box placed off the canvas.
    #
    # inv_summary["overlap_shapes"] is deliberately NOT part of this sum. It is
    # raw pairwise geometry with no notion of z-order or shape kind, so every
    # tinted card counts its own text as an overlap. check_semantic_overlaps owns
    # the overlap verdict; this number is carried in the report for transparency.
    layout_defects = inv_summary["overflow_shapes"]
    if not inv_summary["available"]:
        reasons.append("layout: the layout inventory produced no report")
    if layout_defects:
        reasons.append(
            f"layout defects: {layout_defects} shapes "
            f"({inv_summary['slide_overflow_shapes']} slide-overflow, "
            f"{inv_summary['frame_overflow_shapes']} frame-overflow, "
            f"{inv_summary['off_slide_shapes']} off-slide)"
        )
    if not contrast["pass"]:
        reasons.append(f"contrast: {len(contrast['violations'])} run(s) below WCAG floor")
    if not overlaps["pass"]:
        reasons.append(
            f"overlap: {len(overlaps['violations'])} picture/table shape(s) drawn over other content"
        )
    if not semantic["pass"]:
        reasons.append(
            f"semantic overlap: {len(semantic['violations'])} content pair(s) occluding each other"
        )
    if not objective["pass"]:
        failed = [name for name, check in objective.items() if isinstance(check, dict) and not check.get("pass", True)]
        reasons.append("objective quality: " + ", ".join(failed))

    return {
        "file": str(pptx),
        "pass": not reasons,
        "integrity": {"pass": integrity["pass"], "problems": integrity["problems"][:20]},
        "craft": {"pass": craft["pass"], "slides": craft["slides"], "findings": craft["findings"][:40]},
        "craft_floor": {"pass": not extras_high,
                        "findings": [f for f in extras["findings"] if f["severity"] == "HIGH"][:40],
                        "advisories": [f for f in extras["findings"] if f["severity"] != "HIGH"][:40]},
        "failure_reasons": reasons,
        "anti_slop": {
            "pass": validate_report.get("pass"),
            "slop_score": validate_report.get("slop_score"),
            "review_warnings": validate_report.get("checks", {}).get("forbidden_terms", {}).get("soft_warnings", []),
        },
        "layout": {**inv_summary, "picture_table_overlaps": len(overlaps["violations"])},
        "contrast": contrast,
        "overlaps": overlaps,
        "semantic_overlaps": semantic,
        "objective": objective,
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Honest renderer-agnostic QA gate for PPTX decks.")
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--sibling", type=Path, help="the same source compiled under another tonality (OF-116)")
    args = parser.parse_args(argv)
    report = qa(args.pptx.resolve(), args.sibling.resolve() if args.sibling else None)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["pass"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
