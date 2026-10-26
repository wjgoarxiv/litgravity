#!/usr/bin/env python3
"""craft_extras.py — the office subset of the family craft floor, for the lit-pptx QA gate.

deck_craft.py answers "does the deck look finished"; these checks answer "does it carry the
reflex defaults a designer would flag", each as a number a gate can hold:

  OF-101  at most four accent hue families per slide (among solid shape fills); five fail
  OF-102  body text-box measure (characters per wrapped line, CJK or Latin)
  OF-103  numeric table columns aligned right
  OF-104  concentric radius on a nested rounded frame
  OF-105  gap between cards at least twice the gap between shapes inside a card
  OF-106  no gradient-filled text
  OF-107  no glow effect
  OF-108  no emoji bullets
  OF-109  no tall empty band below the last content block, on the slide or in a card
  OF-110  a deck of 8+ slides mixes at least three title treatments; each drawn as it is named
  OF-111  content slides vary their composition (title zone x body partition x dominant content)
  OF-112  the median empty band under the body's last block stays at or under 0.20
  OF-113  every content title sits on the shared title frame, every display title on the display frame
  OF-114  titles and the cover subtitle are topic labels, not declarative sentences
  OF-115  no region left empty: a takeaway column as long as the visual beside it (its values table
          included), a title panel as tall as its text and no plate under an empty page
  OF-116  (with --sibling) another tonality of the same source sets its titles and bodies in the same places
  OF-117  no light figure card on a dark ground
  OF-118  a bold run-in label is followed by its separator ("요약: …", "Summary: …")
  OF-119  an agenda's title carries no count numeral beside it (its rows are numbered)
  OF-120  a hard offset shadow (no blur) only in a pack that allows one, and there every shadow is hard

A tonality pack may widen two of these limits through its own `qa:` block (listed in
references/anti-slop-checklist.md): `accent-families` and `display-accent-families` (how many accent
hue families a content or display slide may carry before OF-101 fails, five or six; the default is
four), and `offset-shadow` (OF-120).
The gate finds the pack through the deck's tonality stamp; a deck without a stamp keeps the defaults.

Every threshold lives in THRESHOLDS. A HIGH finding fails the gate; a MEDIUM finding is
advisory and is reported, not blocking. `tier` is "measured" when read off the OOXML and
"derived" when it rests on an inherited value (a table style's alignment).

Usage:
    python3 craft_extras.py deck.pptx          # JSON report; exit 0 always
"""

from __future__ import annotations

import os as _os
import sys as _sys

_sys.dont_write_bytecode = True  # the installed skill directory stays read-only
_HERE = _os.path.dirname(_os.path.abspath(__file__))
_sys.path.insert(0, _HERE)
from office_runtime_bootstrap import ensure_runtime  # noqa: E402

ensure_runtime(["pptx"])

import colorsys  # noqa: E402
import json  # noqa: E402
import re  # noqa: E402
from pathlib import Path  # noqa: E402

from pptx import Presentation  # noqa: E402
from pptx.enum.dml import MSO_FILL  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE_TYPE  # noqa: E402
from pptx.enum.text import PP_ALIGN  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402

import deck_craft  # noqa: E402
from layout_inventory import DEFAULT_INSETS, EMU_PER_IN, PT_PER_IN, _attr_int, paragraph_runs, wrap_segment  # noqa: E402

THRESHOLDS = {
    "accent": {"min_lightness": 0.12, "max_lightness": 0.93, "min_saturation": 0.12, "family_deg": 20.0,
               "content_high": 5, "content_medium": 5, "display_high": 5, "display_medium": 5},
    "measure": {"latin_max": 90, "cjk_max": 38, "cjk_share": 0.5, "narrow_min": 10, "narrow_lines": 3},
    "numeric": {"share": 0.7, "min_rows": 2},
    "radius": {"inset_spread_in": 0.06, "tolerance_in": 0.02, "min_outer_in": 0.04, "default_adj": 16667},
    # MEDIUM until the family's read-only re-run on its calibration decks promotes it.
    "grouping": {"ratio": 2.0, "severity": "MEDIUM"},
    "band": {"high": 0.35, "medium": 0.25, "row_tolerance_in": 0.05},
}
EMOJI = re.compile("[\U0001F000-\U0001FAFF☀-➿]")
NUMERIC_STRIP = re.compile(r"[\s,.%▲▼+\-−±()$€£¥₩]")


def _finding(check, severity, slide, shape, detail, hint, tier="measured"):
    return {"check": check, "severity": severity, "tier": tier, "slide": slide, "shape": shape, "detail": detail, "hint": hint}


def _hsl(rgb):
    r, g, b = (int(str(rgb)[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, lightness, s = colorsys.rgb_to_hls(r, g, b)
    return h * 360.0, s, lightness


def _families(hues, width):
    """Greedy hue clustering: a hue within `width` degrees of a family joins it."""
    families = []
    for hue in sorted(hues):
        for family in families:
            if min(abs(hue - family), 360 - abs(hue - family)) <= width:
                break
        else:
            families.append(hue)
    return families


def _solid_fill(shape):
    try:
        if shape.fill.type == MSO_FILL.SOLID:
            return shape.fill.fore_color.rgb
    except (AttributeError, TypeError, ValueError, NotImplementedError):
        return None
    return None


QA_KEY = re.compile(r"^  ([a-z-]+):\s*(\S+)\s*$")


def pack_allowances(prs) -> dict:
    """The stamped tonality pack's `qa:` block as {key: int or bool}; empty without a stamp or a block."""
    stamp = STAMP.search(prs.core_properties.subject or "")
    if not stamp or not re.fullmatch(r"[a-z][a-z0-9-]*", stamp.group(1)):
        return {}
    path = Path(_HERE).parent / "templates" / "tonalities" / stamp.group(1) / "pack.yaml"
    if not path.is_file():
        return {}
    allowed, inside = {}, False
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip() or line.lstrip().startswith("#"):
            continue
        if not line.startswith(" "):
            inside = line.rstrip() == "qa:"
            continue
        match = QA_KEY.match(line) if inside else None
        if match:
            value = match.group(2)
            allowed[match.group(1)] = value == "true" if value in ("true", "false") else int(value) if value.isdigit() else value
    return allowed


def check_shadows(number, slide, allowed):
    """OF-120: a hard offset shadow (blur 0) is a pack's look or nothing; a pack that has it keeps every shadow hard."""
    found = []
    for shape in slide.shapes:
        for shadow in shape._element.iter(qn("a:outerShdw")):
            hard = int(shadow.get("blurRad") or 0) == 0
            if hard and not allowed.get("offset-shadow"):
                found.append(_finding("OF-120", "HIGH", number, shape.name, "hard offset shadow in a pack that does not draw one",
                                      "Remove the shadow; only a pack whose qa block allows offset-shadow draws one."))
            elif not hard and allowed.get("offset-shadow"):
                found.append(_finding("OF-120", "HIGH", number, shape.name, "soft shadow beside the pack's hard offset shadows",
                                      "Draw the shadow with no blur at the pack's offset, or remove it."))
    return found


def check_accents(slide, number, shapes, page_area, display, allowed=None):
    t = dict(THRESHOLDS["accent"])
    # A pack's own limit only widens the default: N families pass, N + 1 fail, with no advisory step between.
    for key, high, medium in (("accent-families", "content_high", "content_medium"), ("display-accent-families", "display_high", "display_medium")):
        if isinstance((allowed or {}).get(key), int) and not isinstance(allowed[key], bool):
            t[high] = t[medium] = max(t[high], allowed[key] + 1)
    hues = []
    for shape in shapes:
        box = deck_craft._box(shape)
        if box is None or box[2] * box[3] >= deck_craft.FULL_BLEED * page_area:
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            continue
        rgb = _solid_fill(shape)
        if rgb is None:
            continue
        hue, saturation, lightness = _hsl(rgb)
        if t["min_lightness"] <= lightness <= t["max_lightness"] and saturation >= t["min_saturation"]:
            hues.append(hue)
    count = len(_families(hues, t["family_deg"]))
    high, medium = (t["display_high"], t["display_medium"]) if display else (t["content_high"], t["content_medium"])
    if count >= high:
        return [_finding("OF-101", "HIGH", number, None, f"{count} accent hue families", f"Hold a slide to {high - 1} accent hue families; fold the extra fills into a family the slide already uses.")]
    if count >= medium:
        return [_finding("OF-101", "MEDIUM", number, None, f"{count} accent hue families", "Close to the accent limit: fold an extra fill into a family the slide already uses.")]
    return []


def _hangul_share(text):
    letters = [c for c in text if not c.isspace()]
    return sum(1 for c in letters if "가" <= c <= "힣") / max(1, len(letters))


def _wrapped_lines(shape):
    body_pr = shape.text_frame._txBody.find(qn("a:bodyPr"))
    insets = {k: _attr_int(body_pr, k, v) for k, v in DEFAULT_INSETS.items()}
    wrap = body_pr.get("wrap", "square") if body_pr is not None else "square"
    inner = (shape.width - insets["lIns"] - insets["rIns"]) / EMU_PER_IN * PT_PER_IN
    lines = 0
    for paragraph in shape.text_frame.paragraphs:
        if not paragraph.text.strip():
            continue
        for segment in paragraph_runs(paragraph, 18.0):
            lines += wrap_segment(segment, max(inner, 1.0), wrap)[0]
    return lines


def check_measure(number, shapes, title):
    t = THRESHOLDS["measure"]
    found = []
    for shape in shapes:
        if shape is title or not getattr(shape, "has_text_frame", False) or not shape.has_text_frame or shape.name.startswith("lit-notice"):
            continue
        size = deck_craft._max_pt(shape)
        if not deck_craft.CAPTION_PT < size < deck_craft.DISPLAY_PT:
            continue
        text = "".join(p.text for p in shape.text_frame.paragraphs)
        lines = _wrapped_lines(shape)
        if lines < 2 or not text.strip():
            continue
        per_line = len(text) / lines
        cjk = _hangul_share(text) >= t["cjk_share"]
        ceiling = t["cjk_max"] if cjk else t["latin_max"]
        if per_line > ceiling:
            found.append(_finding("OF-102", "HIGH", number, shape.name, f"{per_line:.0f} chars per line over {lines} lines{' (Korean)' if cjk else ''}",
                                  f"Narrow the text box or split the text; keep body lines at or under {ceiling} characters."))
        elif per_line < t["narrow_min"] and lines >= t["narrow_lines"]:
            found.append(_finding("OF-102", "MEDIUM", number, shape.name, f"{per_line:.0f} chars per line over {lines} lines",
                                  "Widen the column; very short lines read as broken wrapping."))
    return found


def _numeric(text):
    stripped = NUMERIC_STRIP.sub("", text)
    return bool(stripped) and stripped.isdigit()


def check_numeric_columns(number, shapes):
    t = THRESHOLDS["numeric"]
    found = []
    for shape in shapes:
        if not getattr(shape, "has_table", False) or not shape.has_table:
            continue
        rows = list(shape.table.rows)
        data = rows[1:]
        if len(data) < t["min_rows"]:
            continue
        # A second header-like row (no number in it) above numeric rows is out of scope for now.
        if data and not any(_numeric(cell.text) for cell in data[0].cells) and any(_numeric(cell.text) for row in data[1:] for cell in row.cells):
            continue
        for column in range(len(shape.table.columns)):
            cells = [row.cells[column] for row in data]
            filled = [cell for cell in cells if cell.text.strip()]
            numbers = [cell for cell in filled if _numeric(cell.text)]
            if len(numbers) < t["min_rows"] or len(numbers) / max(1, len(filled)) < t["share"]:
                continue
            explicit = [cell for cell in numbers if any(p.alignment not in (None, PP_ALIGN.RIGHT) for p in cell.text_frame.paragraphs if p.text.strip())]
            inherited = [cell for cell in numbers if all(p.alignment is None for p in cell.text_frame.paragraphs if p.text.strip())]
            if explicit:
                found.append(_finding("OF-103", "HIGH", number, shape.name, f"column {column + 1}: {len(explicit)} number cell(s) not right-aligned",
                                      "Right-align numeric columns so the digits line up."))
            elif inherited:
                found.append(_finding("OF-103", "MEDIUM", number, shape.name, f"column {column + 1}: alignment inherited, not verified",
                                      "Set right alignment on numeric cells explicitly.", tier="derived"))
    return found


def _radius_in(shape):
    geometry = shape._element.find(".//" + qn("a:prstGeom"))
    if geometry is None or geometry.get("prst") != "roundRect":
        return None
    adj = THRESHOLDS["radius"]["default_adj"]
    guide = geometry.find(".//" + qn("a:gd"))
    if guide is not None and (guide.get("fmla") or "").startswith("val "):
        adj = int(guide.get("fmla").split()[1])
    return adj / 100000.0 * min(shape.width, shape.height) / EMU_PER_IN


def check_concentric(number, shapes):
    t = THRESHOLDS["radius"]
    rounded = [(shape, deck_craft._box(shape), _radius_in(shape)) for shape in shapes]
    rounded = [(s, b, r) for s, b, r in rounded if b is not None and r is not None]
    found = []
    for inner, ibox, iradius in rounded:
        for outer, obox, oradius in rounded:
            if inner is outer or not deck_craft._contains(obox, ibox, tol=0.0):
                continue
            pads = [ibox[0] - obox[0], ibox[1] - obox[1], obox[0] + obox[2] - ibox[0] - ibox[2], obox[1] + obox[3] - ibox[1] - ibox[3]]
            if max(pads) - min(pads) > t["inset_spread_in"] or oradius <= t["min_outer_in"]:
                continue
            expected = max(0.0, oradius - sum(pads) / 4)
            if abs(iradius - expected) > t["tolerance_in"]:
                found.append(_finding("OF-104", "HIGH", number, inner.name, f"inner radius {iradius:.2f}in, expected {expected:.2f}in",
                                      "Set the inner radius to the outer radius minus the frame's padding."))
    return found


def _gap(a, b):
    """Clear distance between two boxes; 0 when they overlap."""
    dx = max(0.0, b[0] - (a[0] + a[2]), a[0] - (b[0] + b[2]))
    dy = max(0.0, b[1] - (a[1] + a[3]), a[1] - (b[1] + b[3]))
    return max(dx, dy)


def check_grouping(number, shapes, page_area):
    """Cards are filled or outlined rectangles holding other shapes; compare their spacing.

    The gap inside a card is the smallest clear gap between two shapes it encloses; a shape's
    inset to the card's own edge is padding and never counts.
    """
    ratio = THRESHOLDS["grouping"]["ratio"]
    severity = THRESHOLDS["grouping"]["severity"]
    boxes = [(shape, deck_craft._box(shape)) for shape in shapes]
    cards = [(s, b) for s, b in boxes if b and s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and not deck_craft._text(s)
             and (deck_craft._filled(s) or deck_craft._outlined(s)) and 0.5 <= b[2] * b[3] < deck_craft.FULL_BLEED * page_area]
    members = {}
    for card, cbox in cards:
        inside = [b for s, b in boxes if b and s is not card and deck_craft._contains(cbox, b, tol=0.02) and b != cbox]
        if inside:
            members[id(card)] = (card, cbox, inside)
    within = {}
    for key, (card, cbox, inside) in members.items():
        gaps = [_gap(a, b) for i, a in enumerate(inside) for b in inside[i + 1:]]
        gaps = [gap for gap in gaps if gap > 0]
        if gaps:
            within[key] = min(gaps)
    found = []
    ordered = list(members.values())
    for i, (a, abox, _) in enumerate(ordered):
        for b, bbox, _ in ordered[i + 1:]:
            inner = max([within[k] for k in (id(a), id(b)) if k in within], default=None)
            between = _gap(abox, bbox)
            if inner is not None and 0 < between < ratio * inner:
                found.append(_finding("OF-105", severity, number, f"{a.name} / {b.name}", f"{between:.2f}in between cards, {inner:.2f}in between shapes inside",
                                      "Space the cards at least twice the gap between their contents apart, or tighten the contents."))
    return found


def check_run_effects(number, slide):
    found = []
    for shape in slide.shapes:
        element = shape._element
        if element.findall(".//" + qn("a:rPr") + "/" + qn("a:gradFill")):
            found.append(_finding("OF-106", "HIGH", number, shape.name, "gradient-filled text", "Use a solid text colour; emphasise with weight or size."))
        if element.findall(".//" + qn("a:glow")):
            found.append(_finding("OF-107", "HIGH", number, shape.name, "glow effect", "Remove the glow; use a solid fill or a thin neutral shadow."))
        for bullet in element.iter(qn("a:buChar")):
            if EMOJI.search(bullet.get("char") or ""):
                found.append(_finding("OF-108", "HIGH", number, shape.name, f"emoji bullet {bullet.get('char')}", "Use the template's bullet or a plain symbol."))
    return found


def _band(top, bottom, boxes):
    """(largest gap before the trailing one, trailing gap) as fractions of the height."""
    height = bottom - top
    if height <= 0:
        return 0.0, 0.0
    edges = [top]
    cursor = top
    for box in sorted(boxes, key=lambda b: b[1]):
        edges.append(max(cursor, box[1]))
        cursor = max(cursor, box[1] + box[3])
        edges.append(cursor)
    edges.append(bottom)
    gaps = [edges[i + 1] - edges[i] for i in range(0, len(edges) - 1, 2)]
    return max(gaps[:-1], default=0.0) / height, gaps[-1] / height


def check_bands(number, shapes, area, boxes, page_area, W, H):
    t = THRESHOLDS["band"]
    found = []
    regions = [("slide", area, boxes)]
    for shape in shapes:
        box = deck_craft._box(shape)
        if not box or shape.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE or deck_craft._text(shape) or not deck_craft._filled(shape):
            continue
        if not 0.5 <= box[2] * box[3] < deck_craft.FULL_BLEED * page_area:
            continue
        if deck_craft._bleeds(box, W, H):
            continue
        inside = [(s, b) for s in shapes for b in [deck_craft._box(s)] if b and s is not shape and deck_craft._contains(box, b, tol=0.02) and b != box]
        if not inside:
            continue
        # A card that is itself display-style (one or two large lines, no visual) is exempt.
        large = [s for s, _ in inside if deck_craft._text(s) and deck_craft._max_pt(s) >= deck_craft.DISPLAY_PT]
        if large and len(large) <= 2 and not any(deck_craft._is_table(s) or deck_craft._is_chart(s) for s, _ in inside):
            continue
        regions.append((shape.name, box, [b for _, b in inside]))
    bands = {}
    for name, region, contents in regions:
        if contents:
            bands[name] = _band(region[1], region[1] + region[3], [b for b in contents if b[1] < region[1] + region[3]])
    # Cards sharing a top edge and height form a row; the row's height imposes a trailing band on
    # the card with the least text, so the row counts its smallest trailing band.
    tol = t["row_tolerance_in"]
    cards = [(name, region) for name, region, _ in regions if name != "slide" and name in bands]
    for name, region in cards:
        row = [other for other, box in cards if abs(box[1] - region[1]) <= tol and abs(box[3] - region[3]) <= tol]
        if len(row) > 1:
            bands[name] = (bands[name][0], min(bands[other][1] for other in row))
    for name, region, contents in regions:
        if name not in bands:
            continue
        fraction = max(bands[name])
        if fraction > t["high"]:
            found.append(_finding("OF-109", "HIGH", number, None if name == "slide" else name, f"{fraction:.0%} of the {'content area' if name == 'slide' else 'card'} is one empty band",
                                  "Fill the band with the message's visual, shrink the region to its content, or merge slides."))
        elif fraction > t["medium"]:
            found.append(_finding("OF-109", "MEDIUM", number, None if name == "slide" else name, f"{fraction:.0%} empty band",
                                  "Tighten the region around its content."))
    return found


# ── Deck output checks (OF-110..OF-113) ─────────────────────────────────────
#
# These read the compiled deck as a whole: which title treatment each slide really got, how the
# body is partitioned, and how much of the body is left empty. The treatment comes from the title
# frame's geometry and its companions; the name the engine writes (`title@<treatment>`) is only
# compared with it, never trusted on its own. A deck the engine did not build (no `title@` name)
# is out of scope. Decks on a legacy template report the variety checks and the deck median band
# as MEDIUM: whoever names a legacy template has asked for its single look. All geometry is in points.

TREATMENTS = ("top-rule", "top-plain-large", "side-rail", "band", "statement", "overlay", "kicker-numeral", "bottom-anchor")
# Every content treatment styles the one shared title frame (slide-grid.js), so a slide's title "zone" in the
# composition signature is its treatment: the look varies, the place does not.
ZONE_OF = {t: t for t in TREATMENTS}
OUTPUT = {
    "tol": 2.0, "size_tol": 0.5,
    "min_slides": 8, "min_treatments": 3,
    "variety_cap": 5, "max_share": 0.40,
    "band_floor": 486.0, "footer_top": 492.0, "band_deck_max": 0.20,
    # Per-slide band caps by density (1-2, 3-4, 5-6, 7-8, 9-10); a slide over its cap is MEDIUM.
    "band_caps": ((2, 0.28), (4, 0.24), (6, 0.20), (8, 0.16), (10, 0.12)),
    "anchor_tol": 1.0,
}
DIGITS = re.compile(r"\d+(?:[.:/-]\d+)?\.?")
STAMP = re.compile(r"\blit-pptx tonality=(\S+) density=(\d+) grid=(\w+) variance=(\d+)")
MARGIN = {960: {"airy": 60.0, "standard": 48.0, "dense": 36.0, "compact": 24.0}, 720: {"airy": 54.0, "standard": 42.0, "dense": 30.0, "compact": 24.0}}


def _pt_box(shape):
    try:
        if shape.left is None or shape.width is None:
            return None
        return (shape.left / 12700.0, shape.top / 12700.0, shape.width / 12700.0, shape.height / 12700.0)
    except (TypeError, ValueError):
        return None


def _run_sizes(shape):
    sizes = []
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for paragraph in shape.text_frame.paragraphs:
            sizes += [run.font.size.pt for run in paragraph.runs if run.text.strip() and run.font.size is not None]
    return sizes


def _elements(slide, W, H):
    """Every drawn shape as a dict in points, with its kind and its role (content, card, decoration)."""
    out = []
    for shape in slide.shapes:
        box = _pt_box(shape)
        # The engine names the first shape it draws `family@<id>`: a band or a rule keeps its role;
        # only an empty marker (no width and no height) is skipped. The notice tag and a pack's grid marks
        # in the top margin are page chrome, not composition.
        if box is None or box[2] + box[3] <= 0 or shape.name.startswith(("lit-notice", "grid-mark")):
            continue
        text = deck_craft._text(shape) if getattr(shape, "has_text_frame", False) and shape.has_text_frame else ""
        kind = ("picture" if shape.shape_type == MSO_SHAPE_TYPE.PICTURE else "table" if deck_craft._is_table(shape)
                else "chart" if deck_craft._is_chart(shape) else "text" if text else "shape")
        if kind == "table":
            box = (box[0], box[1], box[2], max(box[3], sum(row.height for row in shape.table.rows) / 12700.0))
        out.append({"shape": shape, "name": shape.name, "kind": kind, "x": box[0], "y": box[1], "w": box[2], "h": box[3],
                    "text": text, "sizes": _run_sizes(shape), "filled": kind == "shape" and deck_craft._filled(shape)})
    for e in out:
        if e["kind"] != "shape":
            e["role"] = "content"
            continue
        bleeds = deck_craft._bleeds((e["x"] / 72, e["y"] / 72, e["w"] / 72, e["h"] / 72), W / 72, H / 72)
        holds = any(o["kind"] != "shape" and _centre_in(o, e) for o in out if o is not e)
        card = e["filled"] and holds and not bleeds and e["w"] >= 100 and e["h"] >= 40 and e["w"] * e["h"] < 0.5 * W * H
        e["role"] = "card" if card else "decoration"
    return out


def _centre_in(e, box, tol=1.0):
    cx, cy = e["x"] + e["w"] / 2, e["y"] + e["h"] / 2
    return box["x"] - tol <= cx <= box["x"] + box["w"] + tol and box["y"] - tol <= cy <= box["y"] + box["h"] + tol


def _footer(e):
    return e["y"] >= OUTPUT["footer_top"]


def classify_title(title, els, W, H):
    """The title treatment a slide shows, by the tests of the design spec in order; first match wins."""
    tol, st = OUTPUT["tol"], OUTPUT["size_tol"]
    size = max(title["sizes"] or [0.0])
    bottom = title["y"] + title["h"]
    others = [e for e in els if e is not title]
    body = [e for e in others if not _footer(e) and e["role"] != "decoration"]

    def behind(e):
        return (e["x"] - tol <= title["x"] and title["x"] + title["w"] <= e["x"] + e["w"] + tol
                and e["y"] - tol <= title["y"] and bottom <= e["y"] + e["h"] + tol)

    if (any(e["kind"] == "picture" and e["w"] * e["h"] >= 0.8 * W * H for e in others) and title["y"] >= 300 - tol
            and any(e["kind"] == "shape" and e["filled"] and behind(e) for e in others)):
        return "overlay"
    if any(e["kind"] == "shape" and e["filled"] and e["y"] <= 2 + tol and 96 - tol <= e["h"] <= 120 + tol
           and e["w"] >= 0.9 * W and _centre_in(title, e, tol) for e in others):
        return "band"
    if (size >= 36 - st and 220 - tol <= title["y"] + title["h"] / 2 <= 380 + tol
            and not any(e["kind"] in ("table", "chart", "picture") or e["role"] == "card" for e in body)
            and sum(1 for e in body if e["kind"] == "text") <= 1):
        return "statement"
    bottom_frame = title["y"] + title["h"]
    content_w = W - 2 * title["x"]
    shapes = [e for e in others if e["kind"] == "shape" and e["filled"]]
    # A heavy rule across the columns closing the title frame.
    if any(2 - st <= e["h"] <= 4 + st and e["w"] >= 0.9 * content_w and bottom_frame - tol <= e["y"] <= bottom_frame + 8 for e in shapes):
        return "bottom-anchor"
    # A bar in the left margin, at least as tall as the title frame.
    if any(e["x"] <= 1 and e["w"] <= 14 and e["h"] >= title["h"] - tol and e["y"] <= title["y"] + tol and e["y"] + e["h"] >= bottom_frame - tol
           for e in shapes):
        return "side-rail"
    for e in others:
        if (e["kind"] == "text" and e["sizes"] and DIGITS.fullmatch(e["text"]) and max(e["sizes"]) >= 40 - st
                and e["x"] >= title["x"] + title["w"] - tol and e["y"] >= title["y"] - tol and e["y"] + e["h"] <= bottom_frame + tol):
            return "kicker-numeral"
    if title["y"] <= 48 + tol:
        # A rule is a line (wider than 16 pt and four times its height), never a list-marker dot.
        rule = any(e["kind"] == "shape" and e["h"] <= 8 and e["w"] >= max(16, 4 * e["h"]) and title["y"] < e["y"] <= bottom + 24
                   and e["x"] < title["x"] + title["w"] and e["x"] + e["w"] > title["x"] for e in others)
        if size >= 36 - st and not rule:
            return "top-plain-large"
        if size <= 32 + st and rule:
            return "top-rule"
    return "unclassified"


def _is_number(e):
    """A text run of 44 pt or more that is mostly digits is a number, not text."""
    compact = re.sub(r"\s+", "", e["text"])
    digits = len(re.findall(r"\d", compact))
    return e["kind"] == "text" and bool(e["sizes"]) and max(e["sizes"]) >= 44 - OUTPUT["size_tol"] and digits / max(1, len(compact)) >= 0.5


def composition(title, treatment, els, W, H, margin):
    """(title zone, body partition, dominant content) of one slide.

    Body boxes are clustered by their x-extent: two boxes overlapping by half the narrower one share
    a column. A box wider than 70 % of the body (a takeaway under two columns) never merges columns.
    """
    zone = ZONE_OF.get(treatment, "none")
    if treatment == "statement":
        return (zone, "one", "text")
    body = []
    for e in els:
        if e is title or _footer(e) or e["role"] == "decoration":
            continue
        if e["kind"] == "shape" and title is not None and _centre_in(title, e, 0):
            continue  # the panel that carries the title
        if e["name"].startswith(("numeral@", "header@")):
            continue
        if e["kind"] == "text" and e["h"] <= 16 and e["y"] + e["h"] >= OUTPUT["band_floor"] - 2:
            continue  # a one-line caption or source on the body floor is the slide's foot, not a column
        body.append(e)
    cards = [e for e in body if e["role"] == "card"]
    blocks = [e for e in body if not any(c is not e and _centre_in(e, c) for c in cards)]
    if any(e["kind"] == "picture" and e["w"] * e["h"] >= 0.8 * W * H for e in els):
        partition = "full-bleed"
    elif not blocks:
        partition = "one"
    else:
        narrow = [e for e in blocks if e["w"] < 0.7 * (W - 2 * margin)] or blocks
        columns = []
        for e in sorted(narrow, key=lambda e: e["x"]):
            hits = [c for c in columns if min(c["r"], e["x"] + e["w"]) - max(c["l"], e["x"]) >= 0.5 * min(c["r"] - c["l"], e["w"])]
            if hits:
                keep = hits[0]
                for other in hits[1:]:
                    keep["l"], keep["r"] = min(keep["l"], other["l"]), max(keep["r"], other["r"])
                    keep["items"] += other["items"]
                    columns.remove(other)
                keep["l"], keep["r"] = min(keep["l"], e["x"]), max(keep["r"], e["x"] + e["w"])
                keep["items"].append(e)
            else:
                columns.append({"l": e["x"], "r": e["x"] + e["w"], "items": [e]})
        column_of = {id(item): n for n, c in enumerate(columns) for item in c["items"]}
        rows = []
        for e in sorted(narrow, key=lambda e: e["y"]):
            if rows and e["y"] < rows[-1]["bottom"] - 2:
                rows[-1]["bottom"] = max(rows[-1]["bottom"], e["y"] + e["h"])
                rows[-1]["items"].append(e)
            else:
                rows.append({"bottom": e["y"] + e["h"], "items": [e]})
        grid_rows = sum(1 for r in rows if len({column_of[id(i)] for i in r["items"]}) >= 2)
        k = len(columns)
        if k >= 4 or (k >= 2 and grid_rows >= 2):
            partition = "grid"
        elif k == 3:
            partition = "three"
        elif k == 2:
            a, b = (c["r"] - c["l"] for c in columns)
            partition = "two-even" if min(a, b) / max(a, b) >= 0.85 else "two-asym"
        else:
            partition = "one"
    area = {}
    kinds = {"picture": "image", "table": "table", "chart": "chart"}
    for e in blocks:
        if e["kind"] in kinds:
            key = kinds[e["kind"]]
        elif e["kind"] == "text":
            key = "number" if _is_number(e) else "text"
        else:
            continue
        area[key] = area.get(key, 0.0) + e["w"] * e["h"]
    for c in cards:
        inner = [e for e in body if e is not c and _centre_in(e, c)]
        key = "number" if any(_is_number(e) for e in inner) else "text"
        area[key] = area.get(key, 0.0) + c["w"] * c["h"]
    dominant = max(area, key=area.get) if area else "text"
    return (zone, partition, dominant)


def bottom_band(title, els, top_margin=36.0):
    """Share of the body zone left empty under its lowest content, from the title bottom (or the top
    margin when the title sits low) to the body floor; decoration, footer and notice excluded."""
    floor = OUTPUT["band_floor"]
    top = title["y"] + title["h"] if title is not None and title["y"] + title["h"] < floor * 0.6 else top_margin
    if floor - top < 10:
        return None
    ink = [e for e in els if e["role"] != "decoration" and not _footer(e) and e["y"] < floor and e["y"] + e["h"] > top]
    lowest = max((min(floor, e["y"] + e["h"]) for e in ink), default=top)
    return max(0.0, (floor - lowest) / (floor - top))


def _slide_kind(family):
    for kind in ("cover", "section", "closing"):
        if family == kind or family.startswith(kind + "-"):
            return kind
    return "content"


def _variance_max(variance):
    return 2 if variance <= 3 else 3 if variance == 4 else 4 if variance <= 6 else 5


def read_deck(prs) -> dict:
    """Every slide's elements, family, kind, written and drawn title treatment, and the deck's dials."""
    W, H = prs.slide_width / 12700.0, prs.slide_height / 12700.0
    stamp = STAMP.search(prs.core_properties.subject or "")
    density, variance = (int(stamp.group(2)), int(stamp.group(4))) if stamp else (5, 5)
    margin = MARGIN[960 if round(W) == 960 else 720].get(stamp.group(3) if stamp else "standard", 48.0)
    slides = []
    for number, slide in enumerate(prs.slides, start=1):
        els = _elements(slide, W, H)
        family = next((s.name.split("@", 1)[1] for s in slide.shapes if s.name.startswith("family@")), "")
        title = next((e for e in els if e["name"].startswith("title@")), None)
        written = title["name"].split("@", 1)[1] if title else None
        treatment = classify_title(title, els, W, H) if title and title["sizes"] else None
        s = {"n": number, "W": W, "H": H, "slide": slide, "els": els, "family": family, "kind": _slide_kind(family), "title": title,
             "written": written, "treatment": treatment}
        if s["kind"] == "content":
            s["signature"] = composition(title, treatment, els, W, H, margin)
        slides.append(s)
    return {"W": W, "H": H, "stamp": stamp, "density": density, "variance": variance, "margin": margin, "slides": slides,
            "legacy": any(s["written"] == "legacy" for s in slides), "engine": any(s["written"] for s in slides)}


def analyse_output(prs) -> list:
    """OF-110..OF-113 and OF-115 over the whole deck."""
    deck = read_deck(prs)
    W, H, stamp, density, variance, slides = deck["W"], deck["H"], deck["stamp"], deck["density"], deck["variance"], deck["slides"]
    if not deck["engine"]:
        return []  # not built by this engine: nothing to compare against
    variety = "MEDIUM" if deck["legacy"] else "HIGH"
    found = []
    for s in slides:
        if s["written"] in TREATMENTS and s["treatment"] not in (None, s["written"]):
            found.append(_finding("OF-110", "HIGH", s["n"], s["title"]["name"], f"the title is drawn as {s['treatment']}, named {s['written']}",
                                  "Use the treatment's own geometry and companions; the written name must match what is drawn."))
    # OF-110: distinct title treatments over the slides that carry one. Cover and section titles
    # are variants of their own (a section keeps the statement frame), so they are not counted.
    titled = [s for s in slides if s["kind"] not in ("cover", "section") and s["treatment"]]
    used = sorted({s["treatment"] for s in titled} - {"unclassified"})
    unclassified = [s["n"] for s in titled if s["treatment"] == "unclassified"]
    if len(slides) >= OUTPUT["min_slides"]:
        if variance <= 3 and len(used) < OUTPUT["min_treatments"]:
            found.append(_finding("OF-110", "MEDIUM", None, None, f"{len(used)} title treatment(s); variety lowered by the user (variance {variance})",
                                  "Variance below 4 is set only on the user's request; say so in the reply."))
        elif len(used) < OUTPUT["min_treatments"]:
            found.append(_finding("OF-110", variety, None, None, f"{len(used)} title treatment(s) ({', '.join(used) or 'none'}) on {len(slides)} slides",
                                  "Map slide roles to the pack's treatments: data under a side rail or a bottom title, sequences under a numeral, claims as statements."))
        elif stamp and len(used) > _variance_max(variance):
            found.append(_finding("OF-110", variety, None, None, f"{len(used)} title treatments, variance {variance} allows {_variance_max(variance)}",
                                  "Lower the number of treatments or raise variance (by at most 2 without asking)."))
    if unclassified:
        found.append(_finding("OF-110", variety, unclassified[0], None, f"no treatment recognised on slide(s) {', '.join(map(str, unclassified))}",
                              "Place the title on one of the eight treatment frames; never position a title by hand."))
    # OF-111: composition signatures over content slides (cover, section and closing excluded).
    content = [s for s in slides if s["kind"] == "content"]
    if content:
        need = min(OUTPUT["variety_cap"], -(-len(content) * 3 // 5))
        problems = []
        for label, key in (("signature", lambda s: s["signature"]), ("layout", lambda s: s["signature"][:2])):
            counts = {}
            for s in content:
                counts[key(s)] = counts.get(key(s), 0) + 1
            top, top_n = max(counts.items(), key=lambda kv: kv[1])
            if len(counts) < need:
                problems.append(f"{label}: {len(counts)} distinct, needs {need}")
            if top_n / len(content) > OUTPUT["max_share"]:
                slides_hit = [s["n"] for s in content if key(s) == top]
                problems.append(f"{label}: {'/'.join(top)} on {top_n} of {len(content)} content slides ({', '.join(map(str, slides_hit))})")
        if problems and len(content) >= 3:
            found.append(_finding("OF-111", variety, None, None, "; ".join(problems),
                                  "Change the family or the title treatment on the repeated slides; let each slide's job pick its shape."))
    # OF-112: the body is used; per slide against the density cap, the deck median against 0.20.
    cap = next(c for upto, c in OUTPUT["band_caps"] if density <= upto)
    bands = []
    for s in slides:
        if s["kind"] in ("cover", "section") or s["family"] in ("statement", "quote", "closing-statement") or s["treatment"] == "statement":
            continue
        band = bottom_band(s["title"], s["els"])
        if band is None:
            continue
        bands.append(band)
        if band > cap:
            found.append(_finding("OF-112", "MEDIUM", s["n"], None, f"{band:.0%} of the body is empty under the last block (cap {cap:.0%} at density {density})",
                                  "Apply the pack's fill order: larger type, a grown visual, wider gaps, or a family that fits the content."))
    if bands:
        bands.sort()
        mid = len(bands) // 2
        median = bands[mid] if len(bands) % 2 else (bands[mid - 1] + bands[mid]) / 2
        if median > OUTPUT["band_deck_max"]:
            # A legacy template has no fill policies: its sparse pages could only be "fixed" with
            # padding, so the median is advisory there and OF-109 still blocks a hollow slide.
            found.append(_finding("OF-112", variety, None, None, f"median empty band {median:.2f} over {len(bands)} slides (limit {OUTPUT['band_deck_max']:.2f})",
                                  "Give the content slides content that reaches the floor: a takeaway under a table, the next step under an ask, a supporting fact under a claim."))
    # OF-113: every content title on one frame and every display title (a statement, a quote, a closing
    # statement) on another, whatever its treatment: a treatment changes the look, never the place. A title
    # over a full-bleed picture (overlay) keeps its own panel. Slides are grouped by the treatment they name,
    # so a title moved so far that it no longer reads as any treatment is still measured.
    groups = {}
    for s in titled:
        if s["written"] in TREATMENTS and s["written"] != "overlay":
            groups.setdefault("display" if s["written"] == "statement" else "content", []).append(s)
    for frame, group in groups.items():
        modes = {}
        for k in ("x", "y", "w", "h"):
            values = [round(s["title"][k]) for s in group]
            modes[k] = max(set(values), key=values.count)
        for s in group:
            off = [f"{k} {s['title'][k]:.0f} vs {modes[k]}" for k in modes if abs(s["title"][k] - modes[k]) > OUTPUT["anchor_tol"]]
            if off:
                found.append(_finding("OF-113", "HIGH", s["n"], s["title"]["name"], f"{s['treatment']} title off the shared {frame} title frame: {', '.join(off)}",
                                      "Every title sits on the slide grid's title zone; fix the source, never nudge a title."))
    found += empty_regions(deck)
    found += agenda_numerals(deck)
    return found


# ── OF-115: no region left empty ────────────────────────────────────────────
#
# The bottom band (OF-112) reads the body as one page. A region can still stand empty beside a full one:
# the takeaway column beside a chart that runs to the floor, the panel that carries a title over a picture,
# and the open page above a plate. Each reading is a share of its own region (40 %).
REGION = {"share": 0.40}


def _largest_gap(top, bottom, spans):
    """The tallest stretch of [top, bottom] that no (y0, y1) span covers."""
    gap, cursor = 0.0, top
    for y0, y1 in sorted(spans):
        if y1 <= cursor:
            continue
        gap = max(gap, max(0.0, min(y0, bottom) - cursor))
        cursor = max(cursor, y1)
    return max(gap, bottom - cursor)


def _short_column(s, content, title, body_floor):
    """A column of a two-column body that stops short of its neighbour: the largest empty band inside it, from
    the body's top down to where the neighbour ends, is over 40 % of the body. A source or note strip that
    closes the column does not hide the band above it (spec-v2 Amendments 8). Beside the figures of a kpi-row or
    big-number slide the column is measured to the body floor (Amendments 9)."""
    tol = OUTPUT["tol"]
    top_of_title = title["y"] + title["h"]
    W, H = s["W"], s["H"]
    fields = [e for e in s["els"] if e["kind"] == "shape" and e["filled"] and e["w"] * e["h"] >= 0.1 * W * H]
    body = [e for e in content if e["y"] >= top_of_title - tol and e["y"] < body_floor and e["w"] > 0 and e["h"] > 0
            and not any(_centre_in(e, f) for f in fields)]
    if len(body) < 2:
        return None
    left_edge, right_edge = min(e["x"] for e in body), max(e["x"] + e["w"] for e in body)
    for cut in sorted({e["x"] + e["w"] for e in body}):
        a = [e for e in body if e["x"] + e["w"] <= cut + tol]
        b = [e for e in body if e["x"] >= cut - tol]
        # A block across most of the body (a note under both columns) belongs to neither column; a narrower block
        # across the cut means there is no column line here.
        across = [e for e in body if e not in a and e not in b]
        if not a or not b or any(e["w"] < 0.7 * (right_edge - left_edge) for e in across):
            continue
        top = min(e["y"] for e in a + b)
        for column, other in ((a, b), (b, a)):
            # A takeaway column beside a chart that closes with the chart's values: the table is read at its drawn extent.
            values = any(e["kind"] == "chart" for e in other) and any(e["kind"] == "text" for e in column) \
                and not any(e["kind"] in ("chart", "picture") for e in column)
            # A text column, or a picture drawn smaller than its column, beside a column that runs on.
            if not any(e["kind"] in ("text", "picture") for e in column) or (any(e["kind"] in ("chart", "table") for e in column) and not values):
                continue
            # Row heads (a matrix's or a comparison's short labels, each level with its row) are not a column to fill.
            if all(e["kind"] == "text" and e["h"] <= 56 and len(re.sub(r"\s+", "", e["text"])) <= 16 for e in column):
                continue
            # Measured whatever the neighbour's length: a neighbour that stops short still marks how far the column should run.
            reach = min(body_floor, max(e["y"] + e["h"] for e in other))
            # A takeaway column beside the figures of a kpi-row or big-number slide runs to the body floor: the figure
            # rows stop where their spacing ends, not where the slide's content does (spec-v2 Amendments 9).
            # The figures' own labels stand exactly level with figures below the first row; they are part of the stack.
            size = lambda side: max([max(e["sizes"]) for e in side if e["kind"] == "text" and e["sizes"]] or [0.0])
            figures = [e for e in other if e["kind"] == "text" and e["sizes"] and max(e["sizes"]) > size(column) + 4]
            labels = any(abs(e["y"] - f["y"]) <= 0.5 and f["y"] > top + tol for e in column for f in figures)
            if s["family"] in ("kpi-row", "big-number") and figures and not labels:
                reach = body_floor
            gap = _largest_gap(top, reach, [(e["y"], e["y"] + e["h"]) for e in column])
            share = gap / max(1.0, body_floor - top)
            if share > REGION["share"]:
                return _finding("OF-115", "HIGH", s["n"], column[0]["name"],
                                f"the {'picture' if any(e['kind'] == 'picture' for e in column) else 'text'} column beside a longer one leaves {share:.0%} of the body empty inside it",
                                "Fill the column by layout: lead size, the source and values in the column, a narrower column with a wider visual, or the points under the visual.")
    return None


def empty_regions(deck) -> list:
    W, H = deck["W"], deck["H"]
    floor = OUTPUT["band_floor"]
    tol = OUTPUT["tol"]
    found = []
    for s in deck["slides"]:
        title, treatment = s["title"], s["treatment"]
        content = [e for e in s["els"] if e is not title and e["role"] != "decoration" and not _footer(e)
                   and not (e["kind"] == "picture" and e["w"] * e["h"] >= 0.8 * W * H)]
        if s["kind"] in ("content", "closing") and title is not None and title["sizes"]:
            # A takeaway column beside a chart, table or picture that runs on below it.
            body_floor = floor
            short = _short_column(s, content, title, body_floor)
            if short:
                found.append(short)
            for v in ([] if short else [e for e in content if e["kind"] in ("chart", "table", "picture")]):
                beside = [e for e in content if e["kind"] == "text" and e["y"] < v["y"] + v["h"] and e["y"] + e["h"] > v["y"]
                          and (e["x"] >= v["x"] + v["w"] - tol or e["x"] + e["w"] <= v["x"] + tol)]
                # Text set on a colour field (a closing's ask field) stands on its own ground, not in a column.
                fields = [e for e in s["els"] if e["kind"] == "shape" and e["filled"] and e["w"] * e["h"] >= 0.1 * W * H]
                beside = [e for e in beside if not any(_centre_in(e, f) for f in fields)]
                if not beside:
                    continue
                left, right = min(e["x"] for e in beside), max(e["x"] + e["w"] for e in beside)
                in_column = lambda e, l=left, r=right: l - tol <= e["x"] + e["w"] / 2 <= r + tol
                in_visual = lambda e, vv=v: vv["x"] - tol <= e["x"] + e["w"] / 2 <= vv["x"] + vv["w"] + tol
                top = min([v["y"]] + [e["y"] for e in beside])
                column = max(e["y"] + e["h"] for e in content if in_column(e))
                visual = min(body_floor, max(e["y"] + e["h"] for e in content if in_visual(e)))
                share = (visual - column) / max(1.0, body_floor - top)
                if share > REGION["share"]:
                    found.append(_finding("OF-115", "HIGH", s["n"], v["name"], f"the column beside the {v['kind']} leaves {share:.0%} of the body empty under its last block",
                                          "Close the column with the source and note lines, the chart's values or the takeaways; never stretch the visual alone."))
                    break
        if s["kind"] in ("cover", "section") or treatment in ("statement", "overlay"):
            if title is None:
                continue
            # A title panel over a picture, and a plate under a statement or quote; a typographic cover's
            # plate is a display device of the cover, not a region its text must fill.
            pictured = any(e["kind"] == "picture" and e["w"] * e["h"] >= 0.8 * W * H for e in s["els"])
            for p in [e for e in s["els"] if e["kind"] == "shape" and e["filled"] and _centre_in(title, e) and e["w"] * e["h"] < 0.9 * W * H]:
                inside = [e for e in content + [title] if e is not p and _centre_in(e, p)]
                if p["y"] > tol and pictured:
                    region_floor = min(p["y"] + p["h"], floor)
                    share = _largest_gap(p["y"], region_floor, [(e["y"], e["y"] + e["h"]) for e in inside]) / max(1.0, region_floor - p["y"])
                    if share > REGION["share"]:
                        found.append(_finding("OF-115", "HIGH", s["n"], p["name"], f"{share:.0%} of the panel that carries the title is empty",
                                              "Make the panel as tall as the title and its lines, standing on the page foot."))
                if p["y"] > tol and treatment == "statement" and p["w"] >= 0.9 * W and not any(e["y"] + e["h"] <= p["y"] + tol for e in content):
                    found.append(_finding("OF-115", "HIGH", s["n"], p["name"], "a plate under part of the page leaves the open page above it empty",
                                          "Set the sentence on the whole-page field or on the open page; a plate needs content above it."))
    return found


# ── OF-116: another tonality of the same source sets everything in the same places ─────
#
# Titles, headers and the body sit on the shared slide grid (slide-grid.js), so two tonalities of one source
# may differ in every look but never in where a slide's title frame and its body's top-left corner stand.
# Slides are compared in order; a slide the other deck lacks is not compared.


def placement(prs) -> list:
    """Per slide: the title frame and the top-left corner of what stands under it, in points."""
    out = []
    for s in read_deck(prs)["slides"]:
        t = s["title"]
        if t is None:
            out.append(None)
            continue
        under = [e for e in s["els"] if e is not t and e["role"] != "decoration" and not _footer(e) and e["y"] >= t["y"] + t["h"] - OUTPUT["tol"]
                 and not e["name"].startswith(("header@", "numeral@", "subtitle@"))]
        corner = (min(e["x"] for e in under), min(e["y"] for e in under)) if under and s["kind"] == "content" else None
        out.append({"n": s["n"], "title": (t["x"], t["y"], t["w"], t["h"]), "body": corner})
    return out


def compare_placement(path: Path, sibling: Path) -> list:
    a, b = placement(Presentation(str(path))), placement(Presentation(str(sibling)))
    tol = OUTPUT["anchor_tol"]
    moved = []
    for x, y in zip(a, b):
        if x is None or y is None:
            continue
        if any(abs(p - q) > tol for p, q in zip(x["title"], y["title"])):
            moved.append(f"slide {x['n']} title at {[round(v) for v in x['title']]} vs {[round(v) for v in y['title']]}")
        elif x["body"] and y["body"] and any(abs(p - q) > tol for p, q in zip(x["body"], y["body"])):
            moved.append(f"slide {x['n']} body from {tuple(round(v) for v in x['body'])} vs {tuple(round(v) for v in y['body'])}")
    if not moved:
        return []
    return [_finding("OF-116", "HIGH", None, None, f"{len(moved)} slide(s) placed differently from {Path(sibling).name}: {'; '.join(moved[:4])}",
                     "A tonality changes the look of a zone, never its place: draw the title and the body on the slide grid's zones.")]


# ── OF-117: no light figure card on a dark ground ───────────────────────────
#
# A chart or diagram exported on white, placed on a dark ground, reads as a white card with small grey
# labels. A figure is a picture of mostly near-white pixels; a photograph rarely is.
FIGURE = {"ground_max": 0.25, "near_white": 235, "share": 0.5, "min_area": 0.03}


def light_figures(number, slide, W, H) -> list:
    try:
        rgb = slide.background.fill.fore_color.rgb if slide.background.fill.type == 1 else None
    except (AttributeError, TypeError, ValueError):
        rgb = None
    if rgb is None:
        return []
    r, g, b = (int(str(rgb)[i:i + 2], 16) / 255 for i in (0, 2, 4))
    if 0.2126 * r + 0.7152 * g + 0.0722 * b > FIGURE["ground_max"]:
        return []
    try:
        import io
        from PIL import Image
    except ImportError:
        return []
    found = []
    for shape in slide.shapes:
        if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
            continue
        box = _pt_box(shape)
        if not box or not FIGURE["min_area"] * W * H <= box[2] * box[3] < 0.8 * W * H:
            continue
        with Image.open(io.BytesIO(shape.image.blob)) as im:
            raw = im.convert("RGB").resize((64, 64)).tobytes()
            share = sum(1 for i in range(0, len(raw), 3) if min(raw[i:i + 3]) >= FIGURE["near_white"]) / (len(raw) / 3)
        if share > FIGURE["share"]:
            found.append(_finding("OF-117", "HIGH", number, shape.name, f"a light figure ({share:.0%} near-white) on a dark ground",
                                  "Supply the figure's dark variant beside it (name.dark.png) or draw it as a native chart."))
    return found


# ── OF-118: a run-in label keeps its separator ─────────────────────────────
#
# A point that opens with a bold label and runs on in regular weight reads as one sentence when nothing
# separates the two ("요약 앱에서 …"): the label takes a colon unless it ends in its own separator.
RUN_IN_SEPARATOR = re.compile(r"[:：.)\]?!–—\-]$")


def check_run_ins(number, slide) -> list:
    found = []
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False) or not shape.has_text_frame or shape.name.startswith(("title@", "lit-notice")):
            continue
        for paragraph in shape.text_frame.paragraphs:
            runs = [r for r in paragraph.runs if r.text]
            if len(runs) < 2 or not runs[0].font.bold or runs[1].font.bold:
                continue
            label, rest = runs[0].text, runs[1].text
            if not label.strip() or not re.search(r"\w", rest) or RUN_IN_SEPARATOR.search(label.strip()) or rest.lstrip()[:1] in ":：":
                continue
            found.append(_finding("OF-118", "HIGH", number, shape.name, f"bold run-in label without a separator: {(label + rest)[:40]}",
                                  "Write the label with its colon (**요약:** …), or let the engine set it; never run a bold label into the sentence."))
    return found


# ── OF-119: no count numeral beside an agenda title ────────────────────────


def agenda_numerals(deck) -> list:
    return [_finding("OF-119", "HIGH", s["n"], s["title"]["name"], "a numeral beside the agenda title repeats the count of its numbered rows",
                     "An agenda takes a title without a numeral; its rows carry the numbers.")
            for s in deck["slides"] if s["family"] == "agenda" and s["treatment"] == "kicker-numeral"]


# ── OF-114: titles are labels ───────────────────────────────────────────────
#
# A title names what the slide covers; the claim goes in the body. Korean is read on its last word
# (어절): a sentence ending or a nominalised claim makes the title a sentence. English is read for a
# final period or a finite verb (an auxiliary, or a report verb in the past or third person). A
# trailing parenthesis is set aside first ("… 요청 (18개월)"). A question is not a claim.

KO_SENTENCE = re.compile(r"(?:니다|[어아해세에예네지군래게]요|죠|[었았였했겠됐](?:다|음)|[가-힣]다|[가-힣](?:함|됨)|(?:있|없)음)$")
KO_NOT_SENTENCE = {"바다", "판다", "소다", "람다", "캐나다", "어젠다", "아젠다", "포함", "보다", "함", "다"}
EN_AUX = re.compile(r"\b(?:is|are|was|were|has|have|had|will|would|can|could|should|must|does|did|won't|isn't|aren't|wasn't|doesn't|didn't)\b", re.I)
EN_VERBS = {
    "grew", "grows", "rose", "rises", "fell", "falls", "doubled", "doubles", "tripled", "triples", "halved", "halves",
    "increased", "increases", "decreased", "decreases", "declined", "declines", "improved", "improves", "dropped", "drops",
    "climbed", "climbs", "beats", "outperforms", "outperformed", "leads", "drives", "drove", "shows", "showed", "needs",
    "makes", "reached", "reaches", "exceeded", "exceeds", "missed", "misses", "stays", "stayed", "remains", "remained",
    "comes", "came", "takes", "took", "works", "worked", "pays", "paid", "saves", "saved", "explains", "explained",
    "matters", "wins", "won", "lost", "loses", "cuts", "lifts", "lifted", "slowed", "slows", "fails", "failed",
}


def declarative_title(text: str) -> bool:
    """True when a title or subtitle states a claim as a sentence instead of naming a topic."""
    t = re.sub(r"\s+", " ", str(text or "")).strip()
    t = re.sub(r"\s*\([^()]*\)$", "", t).strip().rstrip("\"'”’")
    if not t or t.endswith("?") or t[0] in "“\"「『‘'":
        return False  # a question, or a quotation set as the slide's statement
    if re.search(r"[가-힣]", t):
        last = re.sub(r"[.!…]+$", "", t.split(" ")[-1])
        return last not in KO_NOT_SENTENCE and bool(KO_SENTENCE.search(last))
    if t.endswith(".") or t.endswith("!"):
        return True
    words = re.findall(r"[A-Za-z']+", t)
    return bool(EN_AUX.search(t)) or any(w.lower() in EN_VERBS for w in words[1:])


def check_labels(prs) -> list:
    """OF-114 over every engine-drawn title and the cover subtitle."""
    found = []
    named = [(number, shape) for number, slide in enumerate(prs.slides, start=1) for shape in slide.shapes
             if shape.name.startswith(("title@", "subtitle@")) and getattr(shape, "has_text_frame", False) and shape.has_text_frame]
    legacy = any(shape.name == "title@legacy" for _, shape in named)
    for number, shape in named:
        text = re.sub(r"\s+", " ", shape.text_frame.text).strip()
        if declarative_title(text):
            what = "subtitle" if shape.name.startswith("subtitle@") else "title"
            found.append(_finding("OF-114", "MEDIUM" if legacy else "HIGH", number, shape.name,
                                  f"the {what} is a sentence: {text[:60]}",
                                  "Name the topic as a label (분기별 매출 추이; Conversion by feed rate) and move the claim into the body: the first takeaway, point or caption."))
    return found


def analyse(path: Path) -> dict:
    prs = Presentation(str(path))
    W, H = prs.slide_width / EMU_PER_IN, prs.slide_height / EMU_PER_IN
    page_area = W * H
    findings = analyse_output(prs) + check_labels(prs)
    allowed = pack_allowances(prs)
    for number, slide in enumerate(prs.slides, start=1):
        shapes = [s for s in slide.shapes if not s.name.startswith("lit-notice")]
        layout = deck_craft.slide_layout(shapes, W, H, page_area)
        findings += check_run_effects(number, slide)
        findings += check_run_ins(number, slide)
        findings += light_figures(number, slide, W * PT_PER_IN, H * PT_PER_IN)
        findings += check_accents(slide, number, shapes, page_area, layout["display"], allowed)
        findings += check_shadows(number, slide, allowed)
        findings += check_numeric_columns(number, shapes)
        findings += check_concentric(number, shapes)
        if layout["display"]:
            continue
        findings += check_measure(number, shapes, layout["title"])
        findings += check_grouping(number, shapes, page_area)
        findings += check_bands(number, shapes, layout["area"], layout["boxes"], page_area, W, H)
    return {"file": str(path), "pass": not any(f["severity"] == "HIGH" for f in findings), "findings": findings}


def main(argv):
    if not argv:
        print(__doc__)
        return 2
    print(json.dumps(analyse(Path(argv[0])), ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main(_sys.argv[1:]))
